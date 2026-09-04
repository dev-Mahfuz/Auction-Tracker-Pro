
import sys
import json
import uuid
import shutil
from pathlib import Path
from datetime import datetime

from PySide6.QtCore import Qt, QSize, Signal
from PySide6.QtGui import QAction, QFont, QIcon, QColor, QShortcut, QKeySequence
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QGridLayout,
    QLabel, QPushButton, QListWidget, QListWidgetItem, QStackedWidget, QFrame,
    QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView, QLineEdit,
    QComboBox, QSpinBox, QDoubleSpinBox, QFileDialog, QMessageBox, QDialog,
    QDialogButtonBox, QFormLayout, QInputDialog, QCheckBox, QTextEdit, QGroupBox,
    QProgressBar, QSplitter
)

from openpyxl import load_workbook, Workbook


APP_NAME = "Auction Tracker Pro"
APP_VERSION = "3.5.0"

# Branding 
DEVELOPER_NAME = "Mahfuz Rahman"
FACEBOOK_URL = "https://www.facebook.com/mahfuzrahmannn"
INSTAGRAM_URL = "https://www.instagram.com/mahfuz_rahmannn/"
GITHUB_URL = "https://github.com/dev-Mahfuz"


DARK_QSS = """
QWidget {
    background: #0f1115;
    color: #e8eaf0;
    font-family: "Segoe UI", "Inter", "Arial";
    font-size: 13px;
}
QMainWindow { background: #0f1115; }
QFrame#Sidebar { background: #14171d; border-right: 1px solid #252a33; }
QFrame#TopBar { background: #11141a; border-bottom: 1px solid #252a33; }
QFrame#Card {
    background: #171b22;
    border: 1px solid #282e39;
    border-radius: 10px;
}
QLabel#Brand { font-size: 18px; font-weight: 700; color: #ffffff; }
QLabel#Muted { color: #8d96a5; }
QLabel#MetricTitle { color: #929baa; font-size: 12px; }
QLabel#MetricValue { color: #ffffff; font-size: 23px; font-weight: 700; }
QLabel#SectionTitle { color: #ffffff; font-size: 20px; font-weight: 700; }
QLabel#PlayerName { color: #ffffff; font-size: 28px; font-weight: 800; }
QPushButton {
    background: #1b2029;
    border: 1px solid #303744;
    border-radius: 7px;
    padding: 8px 12px;
}
QPushButton:hover { background: #232a35; border-color: #3b4656; }
QPushButton:pressed { background: #151922; }
QPushButton#Primary {
    background: #2563eb;
    color: white;
    border: 1px solid #2563eb;
    font-weight: 600;
}
QPushButton#Primary:hover { background: #1d4ed8; }
QPushButton#Success {
    background: #15803d;
    color: white;
    border: 1px solid #15803d;
    font-weight: 700;
}
QPushButton#Danger {
    background: #991b1b;
    color: white;
    border: 1px solid #991b1b;
    font-weight: 700;
}
QPushButton#Ghost {
    background: transparent;
    border: 1px solid #303744;
}
QPushButton#IconButton {
    padding: 4px;
    border-radius: 6px;
    background: transparent;
}
QPushButton#IconButton:hover { background: #232a35; }
QListWidget {
    background: transparent;
    border: none;
    outline: none;
}
QListWidget::item {
    padding: 10px 12px;
    margin: 2px 7px;
    border-radius: 7px;
    color: #aeb6c4;
}
QListWidget::item:selected {
    background: #212733;
    color: #ffffff;
}
QListWidget::item:hover { background: #1a1f28; }
QTableWidget {
    background: #141820;
    alternate-background-color: #171c24;
    border: 1px solid #282e39;
    gridline-color: #242a33;
    border-radius: 8px;
}
QTableWidget::item { padding: 6px; }
QTableWidget::item:selected { background: #26344d; }
QHeaderView::section {
    background: #1a1f28;
    color: #c8cfda;
    padding: 8px;
    border: none;
    border-right: 1px solid #2a303b;
    border-bottom: 1px solid #2a303b;
    font-weight: 600;
}
QLineEdit, QComboBox, QSpinBox, QDoubleSpinBox, QTextEdit {
    background: #12161c;
    border: 1px solid #303744;
    border-radius: 7px;
    padding: 7px 9px;
    selection-background-color: #2563eb;
}
QLineEdit:focus, QComboBox:focus, QSpinBox:focus, QDoubleSpinBox:focus, QTextEdit:focus {
    border: 1px solid #3b82f6;
}
QComboBox QAbstractItemView {
    background: #171b22;
    border: 1px solid #303744;
    selection-background-color: #26344d;
}
QProgressBar {
    border: none;
    background: #252a33;
    border-radius: 5px;
    height: 9px;
    text-align: center;
}
QProgressBar::chunk { background: #2563eb; border-radius: 5px; }
QGroupBox {
    border: 1px solid #282e39;
    border-radius: 9px;
    margin-top: 10px;
    padding: 14px 10px 10px 10px;
    font-weight: 600;
}
QGroupBox::title { subcontrol-origin: margin; left: 12px; padding: 0 5px; }
QToolTip { background: #20252e; color: white; border: 1px solid #374151; padding: 5px; }
"""

LIGHT_QSS = """
QWidget {
    background: #f5f7fb;
    color: #20242c;
    font-family: "Segoe UI", "Inter", "Arial";
    font-size: 13px;
}
QMainWindow { background: #f5f7fb; }
QFrame#Sidebar { background: #ffffff; border-right: 1px solid #e2e7ef; }
QFrame#TopBar { background: #ffffff; border-bottom: 1px solid #e2e7ef; }
QFrame#Card {
    background: #ffffff;
    border: 1px solid #e1e6ef;
    border-radius: 10px;
}
QLabel#Brand { font-size: 18px; font-weight: 700; color: #111827; }
QLabel#Muted { color: #697386; }
QLabel#MetricTitle { color: #697386; font-size: 12px; }
QLabel#MetricValue { color: #111827; font-size: 23px; font-weight: 700; }
QLabel#SectionTitle { color: #111827; font-size: 20px; font-weight: 700; }
QLabel#PlayerName { color: #111827; font-size: 28px; font-weight: 800; }
QPushButton {
    background: #ffffff;
    border: 1px solid #d8dee8;
    border-radius: 7px;
    padding: 8px 12px;
}
QPushButton:hover { background: #f1f4f9; border-color: #c8d0dc; }
QPushButton#Primary {
    background: #2563eb;
    color: white;
    border: 1px solid #2563eb;
    font-weight: 600;
}
QPushButton#Primary:hover { background: #1d4ed8; }
QPushButton#Success {
    background: #15803d;
    color: white;
    border: 1px solid #15803d;
    font-weight: 700;
}
QPushButton#Danger {
    background: #b91c1c;
    color: white;
    border: 1px solid #b91c1c;
    font-weight: 700;
}
QPushButton#Ghost { background: transparent; border: 1px solid #d8dee8; }
QPushButton#IconButton { padding: 4px; border-radius: 6px; background: transparent; }
QPushButton#IconButton:hover { background: #eef2f7; }
QListWidget {
    background: transparent;
    border: none;
    outline: none;
}
QListWidget::item {
    padding: 10px 12px;
    margin: 2px 7px;
    border-radius: 7px;
    color: #5d6677;
}
QListWidget::item:selected { background: #e8eefb; color: #1e3a8a; }
QListWidget::item:hover { background: #f1f4f9; }
QTableWidget {
    background: #ffffff;
    alternate-background-color: #fafbfc;
    border: 1px solid #e1e6ef;
    gridline-color: #edf0f5;
    border-radius: 8px;
}
QTableWidget::item { padding: 6px; }
QTableWidget::item:selected { background: #dbe8ff; color: #172554; }
QHeaderView::section {
    background: #f3f5f8;
    color: #4c5566;
    padding: 8px;
    border: none;
    border-right: 1px solid #e2e7ef;
    border-bottom: 1px solid #e2e7ef;
    font-weight: 600;
}
QLineEdit, QComboBox, QSpinBox, QDoubleSpinBox, QTextEdit {
    background: #ffffff;
    border: 1px solid #d8dee8;
    border-radius: 7px;
    padding: 7px 9px;
    selection-background-color: #2563eb;
}
QLineEdit:focus, QComboBox:focus, QSpinBox:focus, QDoubleSpinBox:focus, QTextEdit:focus {
    border: 1px solid #3b82f6;
}
QComboBox QAbstractItemView {
    background: #ffffff;
    border: 1px solid #d8dee8;
    selection-background-color: #dbe8ff;
}
QProgressBar {
    border: none;
    background: #e7ebf1;
    border-radius: 5px;
    height: 9px;
    text-align: center;
}
QProgressBar::chunk { background: #2563eb; border-radius: 5px; }
QGroupBox {
    border: 1px solid #e1e6ef;
    border-radius: 9px;
    margin-top: 10px;
    padding: 14px 10px 10px 10px;
    font-weight: 600;
}
QGroupBox::title { subcontrol-origin: margin; left: 12px; padding: 0 5px; }
"""


COMMON_QSS = """
QSplitter::handle {
    background: rgba(100, 116, 139, 0.18);
}
QSplitter::handle:vertical {
    height: 8px;
    margin: 2px 6px;
    border-radius: 4px;
}
QSplitter::handle:horizontal {
    width: 7px;
    margin: 6px 2px;
    border-radius: 4px;
}
QLabel#CompactHint {
    font-size: 11px;
    color: #718096;
}
"""


PRICE_DISPLAY_MODE = "full"


def set_price_display_mode(mode):
    """UI-only formatting. Stored values always remain full numeric amounts."""
    global PRICE_DISPLAY_MODE
    PRICE_DISPLAY_MODE = "compact" if mode == "compact" else "full"


def compact_number(value):
    value = float(value or 0)
    av = abs(value)
    for threshold, suffix in ((1_000_000_000, "B"), (1_000_000, "M"), (1_000, "K")):
        if av >= threshold:
            n = value / threshold
            if abs(n - round(n)) < 1e-9:
                return f"{int(round(n))}{suffix}"
            return f"{n:.2f}".rstrip("0").rstrip(".") + suffix
    if value.is_integer():
        return f"{int(value):,}"
    return f"{value:,.2f}"


def money(v):
    try:
        value = float(v)
    except Exception:
        value = 0
    if PRICE_DISPLAY_MODE == "compact":
        return compact_number(value)
    if value.is_integer():
        return f"{int(value):,}"
    return f"{value:,.2f}"


def table_item(value, align=None):
    item = QTableWidgetItem(str(value))
    if align is not None:
        item.setTextAlignment(align)
    return item




def resource_path(*parts):
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base.joinpath(*parts)


def icon_button(app, icon_name, tooltip, callback, danger=False):
    btn = QPushButton()
    btn.setObjectName("IconButton")
    btn.setIcon(app.icon(icon_name))
    btn.setIconSize(QSize(19, 19))
    btn.setFixedSize(34, 34)
    btn.setToolTip(tooltip)
    if danger:
        btn.setStyleSheet("QPushButton { border-color: rgba(220,38,38,0.35); }")
    btn.clicked.connect(callback)
    return btn


def action_cell(*buttons):
    w = QWidget()
    w.setStyleSheet("background: transparent;")
    lay = QHBoxLayout(w)
    lay.setContentsMargins(4, 4, 4, 4)
    lay.setSpacing(5)
    lay.addStretch()
    for b in buttons:
        lay.addWidget(b)
    lay.addStretch()
    return w


def polish_table(table, row_height=46):
    """Keep action icons fully visible and make every table easier to scan."""
    table.verticalHeader().setDefaultSectionSize(row_height)
    table.verticalHeader().setMinimumSectionSize(row_height)
    table.horizontalHeader().setMinimumHeight(42)
    table.setWordWrap(False)


class MoneyLineEdit(QLineEdit):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setMinimumHeight(48)
        font = self.font()
        font.setPointSize(16)
        font.setBold(True)
        self.setFont(font)
        self.setAlignment(Qt.AlignRight)
        self.setPlaceholderText("0")
        self.editingFinished.connect(self._format)

    def value(self):
        raw = self.text().replace(",", "").replace(" ", "").strip()
        if not raw:
            return 0.0
        try:
            return max(0.0, float(raw))
        except ValueError:
            return 0.0

    def setValue(self, value):
        value = float(value or 0)
        if value.is_integer():
            self.setText(f"{int(value):,}")
        else:
            self.setText(f"{value:,.2f}")

    def addValue(self, amount):
        self.setValue(self.value() + float(amount))
        self.setFocus()
        self.selectAll()

    def _format(self):
        self.setValue(self.value())

class AuctionData:
    def __init__(self):
        self.file_path = None
        self.data = self.empty_data()

    @staticmethod
    def empty_data():
        return {
            "meta": {
                "name": "Untitled Auction",
                "currency": "৳",
                "budget_per_team": 100000,
                "max_players_per_team": 12,
                "my_team": "",
                "theme": "dark",
                "backup_keep": 10,
                "use_common_base_price": False,
                "common_base_price": 0,
                "price_display_mode": "full",
                "quick_price_profile": "Standard",
                "player_queue": [],
                "live_team_columns": ["team", "players", "count", "spent", "remaining", "slots", "latest_player", "latest_price"],
                "creator_name": "Developer",
                "facebook_url": "https://www.facebook.com/",
                "instagram_url": "https://www.instagram.com/",
                "github_url": "https://github.com/",
                "created_at": datetime.now().isoformat(timespec="seconds"),
            },
            "teams": [],
            "players": [],
            "events": [],
        }

    def new(self, name, teams, budget, max_players, my_team, currency, use_common_base_price=False, common_base_price=0):
        self.file_path = None
        self.data = self.empty_data()
        self.data["meta"].update({
            "name": name.strip() or "Untitled Auction",
            "currency": currency.strip() or "৳",
            "budget_per_team": float(budget),
            "max_players_per_team": int(max_players),
            "my_team": my_team,
            "use_common_base_price": bool(use_common_base_price),
            "common_base_price": float(common_base_price or 0),
            "created_at": datetime.now().isoformat(timespec="seconds"),
        })
        self.data["teams"] = [{"id": str(uuid.uuid4()), "name": t.strip()} for t in teams if t.strip()]
        self.data["players"] = []
        self.data["events"] = []

    def save(self, path=None):
        if path:
            self.file_path = Path(path)
        if not self.file_path:
            raise ValueError("No save path selected.")
        self.clean_queue()
        self.file_path.write_text(json.dumps(self.data, indent=2, ensure_ascii=False), encoding="utf-8")

    def load(self, path):
        raw = json.loads(Path(path).read_text(encoding="utf-8"))
        if "meta" not in raw or "teams" not in raw or "players" not in raw or "events" not in raw:
            raise ValueError("Invalid auction file.")
        raw.setdefault("meta", {})
        defaults = self.empty_data()["meta"]
        for k, v in defaults.items():
            raw["meta"].setdefault(k, v)
        for p in raw.get("players", []):
            p.setdefault("wishlist", False)
            p.setdefault("max_bid", 0)
            p.setdefault("notes", "")
        self.data = raw
        self.file_path = Path(path)
        self.clean_queue()

    @property
    def meta(self): return self.data["meta"]
    @property
    def teams(self): return self.data["teams"]
    @property
    def players(self): return self.data["players"]
    @property
    def events(self): return self.data["events"]

    def effective_base_price(self, value=0):
        if self.meta.get("use_common_base_price", False):
            return float(self.meta.get("common_base_price", 0) or 0)
        return float(value or 0)

    def set_common_base_price(self, enabled, value, apply_existing=True):
        self.meta["use_common_base_price"] = bool(enabled)
        self.meta["common_base_price"] = float(value or 0)
        if enabled and apply_existing:
            for p in self.players:
                p["base_price"] = float(value or 0)

    def set_all_base_prices(self, value):
        value = float(value or 0)
        for p in self.players:
            p["base_price"] = value
        self.meta["common_base_price"] = value

    def add_player(self, name, position="", base_price=0, category="", notes="", wishlist=False, max_bid=0):
        p = {
            "id": str(uuid.uuid4()),
            "name": name.strip(),
            "position": position.strip(),
            "base_price": self.effective_base_price(base_price),
            "category": category.strip(),
            "notes": notes.strip(),
            "wishlist": bool(wishlist),
            "max_bid": float(max_bid or 0),
        }
        self.players.append(p)
        return p

    def update_player(self, pid, **kwargs):
        p = self.get_player(pid)
        if not p:
            return
        if self.meta.get("use_common_base_price", False):
            kwargs["base_price"] = float(self.meta.get("common_base_price", 0) or 0)
        if "max_bid" in kwargs:
            kwargs["max_bid"] = float(kwargs.get("max_bid") or 0)
        p.update(kwargs)

    def delete_player(self, pid):
        if any(e.get("player_id") == pid for e in self.events):
            raise ValueError("This player already has auction history. Undo/delete the related history first.")
        self.data["players"] = [p for p in self.players if p["id"] != pid]
        self.queue_remove(pid)

    def add_team(self, name):
        if not name.strip():
            return
        if any(t["name"].lower() == name.strip().lower() for t in self.teams):
            raise ValueError("Team already exists.")
        self.teams.append({"id": str(uuid.uuid4()), "name": name.strip()})

    def rename_team(self, old_name, new_name):
        new_name = new_name.strip()
        if not new_name:
            raise ValueError("Team name cannot be empty.")
        if old_name == new_name:
            return
        if any(t["name"].lower() == new_name.lower() for t in self.teams if t["name"] != old_name):
            raise ValueError("Another team already uses that name.")
        team = next((t for t in self.teams if t["name"] == old_name), None)
        if not team:
            raise ValueError("Team not found.")
        team["name"] = new_name
        for e in self.events:
            if e.get("type") == "sale" and e.get("team") == old_name:
                e["team"] = new_name
        if self.meta.get("my_team") == old_name:
            self.meta["my_team"] = new_name

    def set_wishlist(self, pid, enabled):
        p = self.get_player(pid)
        if p:
            p["wishlist"] = bool(enabled)

    def wishlist_players(self):
        return [p for p in self.players if p.get("wishlist", False)]

    def delete_team(self, name):
        if name == self.meta.get("my_team"):
            raise ValueError("You cannot delete My Team. Change My Team first.")
        if any(e.get("team") == name for e in self.events if e.get("type") == "sale"):
            raise ValueError("This team already has purchase history.")
        self.data["teams"] = [t for t in self.teams if t["name"] != name]

    def get_player(self, pid):
        return next((p for p in self.players if p["id"] == pid), None)

    def event_for_player(self, pid):
        for e in reversed(self.events):
            if e.get("player_id") == pid:
                return e
        return None

    def player_state(self, pid):
        e = self.event_for_player(pid)
        if not e:
            return {"status": "AVAILABLE", "team": "", "price": 0}
        if e["type"] == "sale":
            return {"status": "SOLD", "team": e.get("team", ""), "price": float(e.get("price", 0))}
        if e["type"] == "unsold":
            return {"status": "UNSOLD", "team": "", "price": 0}
        return {"status": "AVAILABLE", "team": "", "price": 0}

    def record_sale(self, pid, team, price):
        if self.player_state(pid)["status"] != "AVAILABLE":
            raise ValueError("Player is not available.")
        if not any(t["name"] == team for t in self.teams):
            raise ValueError("Invalid team.")
        budget = self.team_summary(team)["remaining"]
        if price > budget:
            raise ValueError(f"{team} has only {self.meta['currency']}{money(budget)} remaining.")
        if self.team_summary(team)["players"] >= self.meta["max_players_per_team"]:
            raise ValueError(f"{team} has no player slots remaining.")
        self.events.append({
            "id": str(uuid.uuid4()), "type": "sale", "player_id": pid,
            "team": team, "price": float(price),
            "time": datetime.now().isoformat(timespec="seconds"),
        })
        self.queue_remove(pid)

    def record_unsold(self, pid):
        if self.player_state(pid)["status"] != "AVAILABLE":
            raise ValueError("Player is not available.")
        self.events.append({
            "id": str(uuid.uuid4()), "type": "unsold", "player_id": pid,
            "time": datetime.now().isoformat(timespec="seconds"),
        })
        self.queue_remove(pid)

    def undo_last(self):
        if self.events:
            return self.events.pop()
        return None

    def delete_event(self, event_id):
        self.data["events"] = [e for e in self.events if e["id"] != event_id]

    def team_summary(self, team):
        purchases = [e for e in self.events if e.get("type") == "sale" and e.get("team") == team]
        spent = sum(float(e.get("price", 0)) for e in purchases)
        count = len(purchases)
        budget = float(self.meta.get("budget_per_team", 0))
        max_players = int(self.meta.get("max_players_per_team", 0))
        return {"players": count, "spent": spent, "remaining": budget - spent, "slots": max_players - count}

    def team_purchases(self, team):
        return [e for e in self.events if e.get("type") == "sale" and e.get("team") == team]

    def available_players(self):
        return [p for p in self.players if self.player_state(p["id"])["status"] == "AVAILABLE"]

    def safe_max_bid(self, team):
        summary = self.team_summary(team)
        if summary["slots"] <= 0:
            return 0
        av = self.available_players()
        base_prices = [float(p.get("base_price", 0)) for p in av if float(p.get("base_price", 0)) > 0]
        reserve_floor = min(base_prices) if base_prices else 0
        reserve_for_future_slots = max(0, summary["slots"] - 1) * reserve_floor
        return max(0, summary["remaining"] - reserve_for_future_slots)

    # Player Queue ---------------------------------------------------------
    def clean_queue(self):
        valid = {p["id"] for p in self.available_players()}
        seen = set()
        cleaned = []
        for pid in self.meta.get("player_queue", []):
            if pid in valid and pid not in seen:
                cleaned.append(pid); seen.add(pid)
        self.meta["player_queue"] = cleaned

    def queue_ids(self):
        self.clean_queue()
        return list(self.meta.get("player_queue", []))

    def queue_players(self):
        return [self.get_player(pid) for pid in self.queue_ids() if self.get_player(pid)]

    def queue_add(self, pid):
        if not self.get_player(pid) or self.player_state(pid)["status"] != "AVAILABLE":
            return
        q = self.meta.setdefault("player_queue", [])
        if pid not in q:
            q.append(pid)

    def queue_remove(self, pid):
        q = self.meta.setdefault("player_queue", [])
        self.meta["player_queue"] = [x for x in q if x != pid]

    def queue_move(self, pid, delta):
        q = self.meta.setdefault("player_queue", [])
        if pid not in q:
            return
        i = q.index(pid)
        j = max(0, min(len(q) - 1, i + delta))
        if i != j:
            q[i], q[j] = q[j], q[i]


class NewAuctionDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Create New Auction")
        self.setMinimumWidth(600)
        root = QVBoxLayout(self)

        title = QLabel("Create Auction")
        title.setObjectName("SectionTitle")
        subtitle = QLabel("Set auction rules, teams and optional common base price. No database is used.")
        subtitle.setObjectName("Muted")
        root.addWidget(title); root.addWidget(subtitle)

        form = QFormLayout()
        self.name = QLineEdit("My Auction")
        self.currency = QLineEdit("৳")
        self.budget = QDoubleSpinBox(); self.budget.setRange(0, 1_000_000_000_000); self.budget.setDecimals(0); self.budget.setValue(100000); self.budget.setSingleStep(5000)
        self.max_players = QSpinBox(); self.max_players.setRange(1, 100); self.max_players.setValue(12)
        self.team_count = QSpinBox(); self.team_count.setRange(2, 50); self.team_count.setValue(8)
        self.common_base_enabled = QCheckBox("Use the same base price for every player")
        self.common_base = QDoubleSpinBox(); self.common_base.setRange(0, 1_000_000_000_000); self.common_base.setDecimals(0); self.common_base.setSingleStep(500); self.common_base.setValue(1000); self.common_base.setEnabled(False)
        self.common_base_enabled.toggled.connect(self.common_base.setEnabled)

        form.addRow("Auction name", self.name)
        form.addRow("Currency symbol", self.currency)
        form.addRow("Budget per team", self.budget)
        form.addRow("Max players per team", self.max_players)
        form.addRow("Number of teams", self.team_count)
        form.addRow("Common base price", self.common_base_enabled)
        form.addRow("Base price amount", self.common_base)
        root.addLayout(form)

        root.addWidget(QLabel("Team names (one per line)"))
        self.teams = QTextEdit(); self.teams.setPlaceholderText("Thunder FC\nTigers\nWarriors\nFalcons"); self.teams.setPlainText("\n".join([f"Team {i}" for i in range(1, 9)])); self.teams.setMinimumHeight(145)
        root.addWidget(self.teams)

        myrow = QHBoxLayout(); myrow.addWidget(QLabel("My Team")); self.my_team = QComboBox(); myrow.addWidget(self.my_team, 1); root.addLayout(myrow)
        self.team_count.valueChanged.connect(self.sync_team_lines); self.teams.textChanged.connect(self.refresh_my_team); self.refresh_my_team()

        buttons = QDialogButtonBox(QDialogButtonBox.Cancel | QDialogButtonBox.Ok)
        buttons.button(QDialogButtonBox.Ok).setText("Create Auction")
        buttons.accepted.connect(self.validate_and_accept); buttons.rejected.connect(self.reject); root.addWidget(buttons)

    def sync_team_lines(self):
        count = self.team_count.value(); current = [x.strip() for x in self.teams.toPlainText().splitlines() if x.strip()]
        while len(current) < count: current.append(f"Team {len(current)+1}")
        current = current[:count]
        self.teams.blockSignals(True); self.teams.setPlainText("\n".join(current)); self.teams.blockSignals(False); self.refresh_my_team()

    def refresh_my_team(self):
        old = self.my_team.currentText(); names = [x.strip() for x in self.teams.toPlainText().splitlines() if x.strip()]
        self.my_team.blockSignals(True); self.my_team.clear(); self.my_team.addItems(names)
        if old in names: self.my_team.setCurrentText(old)
        self.my_team.blockSignals(False)

    def validate_and_accept(self):
        teams = [x.strip() for x in self.teams.toPlainText().splitlines() if x.strip()]
        if len(teams) != self.team_count.value():
            QMessageBox.warning(self, "Team Count", f"Please enter exactly {self.team_count.value()} team names."); return
        if len({t.lower() for t in teams}) != len(teams):
            QMessageBox.warning(self, "Duplicate Teams", "Team names must be unique."); return
        if not self.name.text().strip():
            QMessageBox.warning(self, "Auction Name", "Enter an auction name."); return
        self.accept()

    def values(self):
        teams = [x.strip() for x in self.teams.toPlainText().splitlines() if x.strip()]
        return {"name": self.name.text().strip(), "currency": self.currency.text().strip() or "৳", "budget": self.budget.value(),
                "max_players": self.max_players.value(), "teams": teams, "my_team": self.my_team.currentText(),
                "use_common_base_price": self.common_base_enabled.isChecked(), "common_base_price": self.common_base.value()}


class PlayerDialog(QDialog):
    def __init__(self, parent=None, player=None, common_base_enabled=False, common_base_price=0):
        super().__init__(parent)
        self.setWindowTitle("Edit Player" if player else "Add Player")
        self.setMinimumWidth(450)
        form = QFormLayout(self)

        self.name = QLineEdit(player.get("name", "") if player else "")
        self.position = QLineEdit(player.get("position", "") if player else "")
        self.category = QLineEdit(player.get("category", "") if player else "")
        self.base_price = QDoubleSpinBox(); self.base_price.setRange(0, 1_000_000_000_000); self.base_price.setDecimals(0); self.base_price.setSingleStep(500)
        self.base_price.setValue(float(player.get("base_price", common_base_price)) if player else float(common_base_price or 0))
        self.base_price.setEnabled(not common_base_enabled)
        self.max_bid = QDoubleSpinBox(); self.max_bid.setRange(0, 1_000_000_000_000); self.max_bid.setDecimals(0); self.max_bid.setSingleStep(500)
        self.max_bid.setValue(float(player.get("max_bid", 0)) if player else 0)
        self.notes = QTextEdit(player.get("notes", "") if player else ""); self.notes.setMaximumHeight(90)
        self.wishlist = QCheckBox("Add to wishlist / target list"); self.wishlist.setChecked(bool(player.get("wishlist", False)) if player else False)

        form.addRow("Player name", self.name); form.addRow("Position / Role", self.position); form.addRow("Category", self.category)
        form.addRow("Base price", self.base_price); form.addRow("Personal max bid", self.max_bid); form.addRow("Notes", self.notes); form.addRow("", self.wishlist)
        if common_base_enabled:
            note = QLabel(f"Common base price is enabled: {common_base_price:,.0f}"); note.setObjectName("Muted"); form.addRow("", note)

        buttons = QDialogButtonBox(QDialogButtonBox.Cancel | QDialogButtonBox.Ok); buttons.accepted.connect(self._accept); buttons.rejected.connect(self.reject); form.addRow(buttons)

    def _accept(self):
        if not self.name.text().strip(): QMessageBox.warning(self, "Player Name", "Player name is required."); return
        self.accept()

    def values(self):
        return {"name": self.name.text().strip(), "position": self.position.text().strip(), "category": self.category.text().strip(),
                "base_price": self.base_price.value(), "max_bid": self.max_bid.value(), "notes": self.notes.toPlainText().strip(), "wishlist": self.wishlist.isChecked()}


class SaleEditDialog(QDialog):
    def __init__(self, teams, player_name, event, currency, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Edit Sale")
        self.setMinimumWidth(400)
        form = QFormLayout(self)

        label = QLabel(player_name)
        label.setObjectName("SectionTitle")
        form.addRow("Player", label)

        self.team = QComboBox()
        self.team.addItems(teams)
        self.team.setCurrentText(event.get("team", ""))

        self.price = MoneyLineEdit()
        self.price.setValue(float(event.get("price", 0)))
        price_wrap = QWidget()
        price_lay = QHBoxLayout(price_wrap)
        price_lay.setContentsMargins(0, 0, 0, 0)
        price_lay.addWidget(QLabel(currency))
        price_lay.addWidget(self.price, 1)

        form.addRow("Team", self.team)
        form.addRow("Sold price", price_wrap)

        buttons = QDialogButtonBox(QDialogButtonBox.Cancel | QDialogButtonBox.Ok)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)


class MetricCard(QFrame):
    def __init__(self, title, value="0"):
        super().__init__()
        self.setObjectName("Card")
        lay = QVBoxLayout(self)
        lay.setContentsMargins(16, 14, 16, 14)
        t = QLabel(title)
        t.setObjectName("MetricTitle")
        self.value = QLabel(value)
        self.value.setObjectName("MetricValue")
        lay.addWidget(t)
        lay.addWidget(self.value)


class OverviewPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 22, 24, 22)
        root.setSpacing(16)

        title = QLabel("Auction Overview")
        title.setObjectName("SectionTitle")
        self.subtitle = QLabel()
        self.subtitle.setObjectName("Muted")
        root.addWidget(title)
        root.addWidget(self.subtitle)

        cards = QGridLayout()
        self.total = MetricCard("Total Players")
        self.sold = MetricCard("Sold")
        self.available = MetricCard("Available")
        self.unsold = MetricCard("Unsold")
        self.my_remaining = MetricCard("My Budget Remaining")
        self.my_slots = MetricCard("My Slots Left")
        for i, c in enumerate([self.total, self.sold, self.available, self.unsold, self.my_remaining, self.my_slots]):
            cards.addWidget(c, i // 3, i % 3)
        root.addLayout(cards)

        row = QHBoxLayout()
        label = QLabel("All Teams")
        label.setObjectName("SectionTitle")
        row.addWidget(label)
        row.addStretch()
        root.addLayout(row)

        self.table = QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels(["Team", "Players", "Spent", "Remaining", "Slots Left"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(True)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        root.addWidget(self.table, 1)

    def refresh(self):
        d = self.app.model
        cur = d.meta.get("currency", "৳")
        states = [d.player_state(p["id"])["status"] for p in d.players]
        sold = states.count("SOLD")
        unsold = states.count("UNSOLD")
        av = states.count("AVAILABLE")
        self.subtitle.setText(f"{d.meta.get('name', '')}  •  My Team: {d.meta.get('my_team', '—')}")
        self.total.value.setText(str(len(d.players)))
        self.sold.value.setText(str(sold))
        self.available.value.setText(str(av))
        self.unsold.value.setText(str(unsold))

        my_team = d.meta.get("my_team", "")
        ms = d.team_summary(my_team) if my_team else {"remaining": 0, "slots": 0}
        self.my_remaining.value.setText(f"{cur}{money(ms['remaining'])}")
        self.my_slots.value.setText(str(ms["slots"]))

        self.table.setRowCount(len(d.teams))
        for r, team in enumerate(d.teams):
            name = team["name"]
            s = d.team_summary(name)
            vals = [name, f"{s['players']}/{d.meta['max_players_per_team']}", f"{cur}{money(s['spent'])}",
                    f"{cur}{money(s['remaining'])}", str(s["slots"])]
            for c, v in enumerate(vals):
                item = table_item(v)
                if name == my_team:
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                self.table.setItem(r, c, item)


class LiveTeamColumnsDialog(QDialog):
    COLUMN_OPTIONS = [
        ("team", "Team"),
        ("players", "Players Bought"),
        ("count", "Count"),
        ("latest_player", "Latest Player"),
        ("latest_price", "Latest Price"),
        ("spent", "Spent"),
        ("remaining", "Remaining Balance"),
        ("slots", "Slots Left"),
        ("safe", "Safe Max Bid"),
        ("avg", "Average Price"),
    ]

    def __init__(self, current, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Manage Live Team Table Columns")
        self.setMinimumWidth(390)
        root = QVBoxLayout(self)
        title = QLabel("Choose table columns")
        title.setObjectName("SectionTitle")
        root.addWidget(title)
        hint = QLabel("Checked columns are shown in the Live Tracker team table. Uncheck a column to remove it; check it again to add it back.")
        hint.setWordWrap(True)
        hint.setObjectName("Muted")
        root.addWidget(hint)
        self.boxes = {}
        current = set(current or [])
        for key, label in self.COLUMN_OPTIONS:
            cb = QCheckBox(label)
            cb.setChecked(key in current or key == "team")
            if key == "team":
                cb.setEnabled(False)
                cb.setToolTip("Team column is always visible")
            self.boxes[key] = cb
            root.addWidget(cb)
        buttons = QDialogButtonBox(QDialogButtonBox.Cancel | QDialogButtonBox.Ok)
        buttons.button(QDialogButtonBox.Ok).setText("Apply Columns")
        buttons.accepted.connect(self._accept)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

    def _accept(self):
        if not any(cb.isChecked() for cb in self.boxes.values()):
            QMessageBox.warning(self, "Columns", "Keep at least one column visible.")
            return
        self.accept()

    def selected_columns(self):
        cols = [key for key, _ in self.COLUMN_OPTIONS if self.boxes[key].isChecked()]
        if "team" not in cols:
            cols.insert(0, "team")
        return cols


class LiveTrackerPage(QWidget):
    """Fast, uncluttered live-auction workspace with a resizable Teams board below."""

    def __init__(self, app):
        super().__init__()
        self.app = app
        self.selected_pid = None

        root = QVBoxLayout(self)
        root.setContentsMargins(18, 12, 18, 12)
        root.setSpacing(7)

        # Compact header: only information needed while the auction is running.
        header = QHBoxLayout()
        title = QLabel("Live Tracker")
        title.setObjectName("SectionTitle")
        header.addWidget(title)
        self.available_count = QLabel()
        self.available_count.setObjectName("Muted")
        header.addWidget(self.available_count)
        header.addStretch()
        hint = QLabel("Ctrl+F search  •  Ctrl+Enter sold  •  Ctrl+U unsold  •  Ctrl+Z undo")
        hint.setObjectName("CompactHint")
        header.addWidget(hint)
        self.undo_btn = QPushButton("Undo Last")
        self.undo_btn.setIcon(self.app.icon("undo"))
        self.undo_btn.setToolTip("Undo the latest auction result (Ctrl+Z)")
        self.undo_btn.clicked.connect(self.undo_last)
        header.addWidget(self.undo_btn)
        root.addLayout(header)

        # Main vertical splitter. Drag the handle down to give Live Tracker more room,
        # or drag it up to see more of the Teams board.
        self.page_splitter = QSplitter(Qt.Vertical)
        self.page_splitter.setChildrenCollapsible(False)
        self.page_splitter.setHandleWidth(8)

        tracker_area = QWidget()
        tracker_layout = QVBoxLayout(tracker_area)
        tracker_layout.setContentsMargins(0, 0, 0, 0)
        tracker_layout.setSpacing(0)

        work_split = QSplitter(Qt.Horizontal)
        work_split.setChildrenCollapsible(False)
        work_split.setHandleWidth(7)

        # ---------------- Player finder ----------------
        left = QFrame()
        left.setObjectName("Card")
        ll = QVBoxLayout(left)
        ll.setContentsMargins(14, 12, 14, 12)
        ll.setSpacing(7)

        self.search = QLineEdit()
        self.search.setPlaceholderText("Search player name, serial, role or category…")
        self.search.setClearButtonEnabled(True)
        self.search.setMinimumHeight(38)
        self.search.textChanged.connect(self.refresh_player_list)
        ll.addWidget(self.search)

        filters = QHBoxLayout()
        filters.setSpacing(7)
        self.role_filter = QComboBox()
        self.category_filter = QComboBox()
        self.role_filter.currentTextChanged.connect(self.refresh_player_list)
        self.category_filter.currentTextChanged.connect(self.refresh_player_list)
        filters.addWidget(self.role_filter, 1)
        filters.addWidget(self.category_filter, 1)
        ll.addLayout(filters)

        self.player_table = QTableWidget(0, 6)
        self.player_table.setHorizontalHeaderLabels(["#", "Player", "Role", "Base", "My Max", "Target"])
        self.player_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        self.player_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.player_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        for c in (3, 4, 5):
            self.player_table.horizontalHeader().setSectionResizeMode(c, QHeaderView.ResizeToContents)
        self.player_table.verticalHeader().setVisible(False)
        self.player_table.verticalHeader().setDefaultSectionSize(36)
        self.player_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.player_table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.player_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.player_table.setAlternatingRowColors(True)
        self.player_table.itemSelectionChanged.connect(self.player_selected)
        self.player_table.setMinimumHeight(155)
        ll.addWidget(self.player_table, 1)

        # Selected-player summary is deliberately compact.
        selected = QFrame()
        selected.setObjectName("Card")
        si = QHBoxLayout(selected)
        si.setContentsMargins(10, 7, 10, 7)
        name_col = QVBoxLayout()
        name_col.setSpacing(1)
        self.player_name = QLabel("No available player")
        self.player_name.setObjectName("Brand")
        self.player_meta = QLabel("Select a player above")
        self.player_meta.setObjectName("Muted")
        name_col.addWidget(self.player_name)
        name_col.addWidget(self.player_meta)
        si.addLayout(name_col, 1)
        self.base_price = QLabel("—")
        self.base_price.setObjectName("MetricValue")
        self.personal_max = QLabel("—")
        self.personal_max.setObjectName("MetricValue")
        for label, value in (("Base", self.base_price), ("My Max", self.personal_max)):
            box = QVBoxLayout()
            box.setSpacing(0)
            small = QLabel(label)
            small.setObjectName("Muted")
            box.addWidget(small)
            box.addWidget(value)
            si.addLayout(box)
        ll.addWidget(selected)
        work_split.addWidget(left)

        # ---------------- Sell panel ----------------
        right = QFrame()
        right.setObjectName("Card")
        rl = QVBoxLayout(right)
        rl.setContentsMargins(16, 12, 16, 12)
        rl.setSpacing(7)
        st = QLabel("Sell Player")
        st.setObjectName("SectionTitle")
        rl.addWidget(st)

        team_label = QLabel("Bought by")
        team_label.setObjectName("Muted")
        rl.addWidget(team_label)
        self.team_combo = QComboBox()
        self.team_combo.setMinimumHeight(40)
        self.team_combo.currentIndexChanged.connect(self.update_budget_info)
        rl.addWidget(self.team_combo)

        self.budget_info = QLabel()
        self.budget_info.setObjectName("Muted")
        self.budget_info.setWordWrap(True)
        rl.addWidget(self.budget_info)

        price_head = QHBoxLayout()
        price_head.addWidget(QLabel("Sold price"))
        price_head.addStretch()
        self.safe_bid = QLabel("—")
        self.safe_bid.setObjectName("Brand")
        price_head.addWidget(QLabel("Safe bid:"))
        price_head.addWidget(self.safe_bid)
        rl.addLayout(price_head)

        pr = QHBoxLayout()
        self.currency_label = QLabel("৳")
        self.currency_label.setObjectName("MetricValue")
        self.price = MoneyLineEdit()
        self.price.textChanged.connect(self.price_changed)
        pr.addWidget(self.currency_label)
        pr.addWidget(self.price, 1)
        rl.addLayout(pr)

        # Quick-price profiles keep the sell panel clean while supporting auctions
        # that move by 10, 500, 1K, 100K, 100M, etc.
        profile_row = QHBoxLayout()
        profile_row.setSpacing(6)
        profile_row.addWidget(QLabel("Quick price"))
        self.quick_profile = QComboBox()
        self.quick_profile.addItems(["Fine", "Standard", "Thousands", "Lakhs", "Millions", "Huge"])
        self.quick_profile.setToolTip("Choose the bid increment scale")
        self.quick_profile.currentTextChanged.connect(self.quick_profile_changed)
        profile_row.addWidget(self.quick_profile, 1)
        self.price_format = QComboBox()
        self.price_format.addItem("Full", "full")
        self.price_format.addItem("Compact K/M/B", "compact")
        self.price_format.setToolTip("Display amounts as 100,000 or 100K/1M")
        self.price_format.currentIndexChanged.connect(self.price_format_changed)
        profile_row.addWidget(self.price_format)
        rl.addLayout(profile_row)

        quick = QHBoxLayout()
        quick.setSpacing(6)
        self.quick_buttons = []
        for i in range(4):
            b = QPushButton("+")
            b.setMinimumHeight(34)
            b.clicked.connect(lambda checked=False, idx=i: self.apply_quick_increment(idx))
            self.quick_buttons.append(b)
            quick.addWidget(b, 1)
        rl.addLayout(quick)

        self.limit_warning = QLabel("")
        self.limit_warning.setWordWrap(True)
        self.limit_warning.setMinimumHeight(20)
        rl.addWidget(self.limit_warning)
        rl.addStretch()

        buttons = QHBoxLayout()
        sold = QPushButton("MARK SOLD")
        sold.setObjectName("Success")
        sold.setIcon(self.app.icon("check"))
        sold.setMinimumHeight(46)
        sold.clicked.connect(self.mark_sold)
        unsold = QPushButton("UNSOLD")
        unsold.setObjectName("Danger")
        unsold.setMinimumHeight(46)
        unsold.clicked.connect(self.mark_unsold)
        buttons.addWidget(sold, 1)
        buttons.addWidget(unsold, 1)
        rl.addLayout(buttons)
        work_split.addWidget(right)
        work_split.setStretchFactor(0, 3)
        work_split.setStretchFactor(1, 2)
        work_split.setSizes([760, 460])

        tracker_layout.addWidget(work_split, 1)
        tracker_area.setMinimumHeight(330)
        self.page_splitter.addWidget(tracker_area)

        # Reuse the exact Teams layout beneath the tracker. This replaces the old
        # wide custom status table so there is only one team UI to learn.
        self.embedded_teams = TeamsPage(self.app, embedded=True)
        self.embedded_teams.setMinimumHeight(225)
        self.page_splitter.addWidget(self.embedded_teams)
        self.page_splitter.setStretchFactor(0, 3)
        self.page_splitter.setStretchFactor(1, 2)
        self.page_splitter.setSizes([430, 330])
        root.addWidget(self.page_splitter, 1)

        QShortcut(QKeySequence("Ctrl+F"), self, activated=lambda: (self.search.setFocus(), self.search.selectAll()))
        QShortcut(QKeySequence("Ctrl+Return"), self, activated=self.mark_sold)
        QShortcut(QKeySequence("Ctrl+Enter"), self, activated=self.mark_sold)
        QShortcut(QKeySequence("Ctrl+U"), self, activated=self.mark_unsold)

    QUICK_PRICE_PROFILES = {
        "Fine": [10, 50, 100, 500],
        "Standard": [500, 1_000, 2_000, 5_000],
        "Thousands": [1_000, 5_000, 10_000, 50_000],
        "Lakhs": [100_000, 200_000, 500_000, 1_000_000],
        "Millions": [1_000_000, 5_000_000, 10_000_000, 50_000_000],
        "Huge": [100_000_000, 200_000_000, 500_000_000, 1_000_000_000],
    }

    @staticmethod
    def increment_label(amount):
        amount = float(amount)
        if amount >= 1_000_000_000:
            return f"+{amount/1_000_000_000:g}B"
        if amount >= 1_000_000:
            return f"+{amount/1_000_000:g}M"
        if amount >= 1_000:
            return f"+{amount/1_000:g}K"
        return f"+{amount:g}"

    def current_team(self):
        return self.team_combo.currentData() or ""

    def quick_profile_changed(self, name):
        if name not in self.QUICK_PRICE_PROFILES:
            return
        self.app.model.meta["quick_price_profile"] = name
        self.update_quick_buttons()
        if self.app.model.file_path:
            try:
                self.app.model.save()
            except Exception:
                pass

    def update_quick_buttons(self):
        profile = self.quick_profile.currentText() or "Standard"
        amounts = self.QUICK_PRICE_PROFILES.get(profile, self.QUICK_PRICE_PROFILES["Standard"])
        for button, amount in zip(self.quick_buttons, amounts):
            button.setText(self.increment_label(amount))
            button.setToolTip(f"Add {int(amount):,} to the current sold price")

    def apply_quick_increment(self, index):
        profile = self.quick_profile.currentText() or "Standard"
        amounts = self.QUICK_PRICE_PROFILES.get(profile, self.QUICK_PRICE_PROFILES["Standard"])
        if 0 <= index < len(amounts):
            self.price.addValue(amounts[index])

    def price_format_changed(self, *_):
        mode = self.price_format.currentData() or "full"
        if self.app.model.meta.get("price_display_mode") == mode:
            return
        self.app.model.meta["price_display_mode"] = mode
        set_price_display_mode(mode)
        self.app.after_change(auto_save=True)

    def refresh(self):
        d = self.app.model
        self.currency_label.setText(d.meta.get("currency", "৳"))
        self.quick_profile.blockSignals(True)
        profile = d.meta.get("quick_price_profile", "Standard")
        if profile not in self.QUICK_PRICE_PROFILES:
            profile = "Standard"
        self.quick_profile.setCurrentText(profile)
        self.quick_profile.blockSignals(False)
        self.update_quick_buttons()
        self.price_format.blockSignals(True)
        mode = d.meta.get("price_display_mode", "full")
        idx = self.price_format.findData(mode)
        self.price_format.setCurrentIndex(idx if idx >= 0 else 0)
        self.price_format.blockSignals(False)
        old_role = self.role_filter.currentText()
        old_cat = self.category_filter.currentText()
        roles = sorted({p.get("position", "").strip() for p in d.players if p.get("position", "").strip()})
        cats = sorted({p.get("category", "").strip() for p in d.players if p.get("category", "").strip()})
        self.role_filter.blockSignals(True)
        self.category_filter.blockSignals(True)
        self.role_filter.clear(); self.role_filter.addItem("All Roles"); self.role_filter.addItems(roles)
        self.category_filter.clear(); self.category_filter.addItem("All Categories"); self.category_filter.addItems(cats)
        if old_role in roles: self.role_filter.setCurrentText(old_role)
        if old_cat in cats: self.category_filter.setCurrentText(old_cat)
        self.role_filter.blockSignals(False)
        self.category_filter.blockSignals(False)

        old_team = self.current_team()
        self.team_combo.blockSignals(True)
        self.team_combo.clear()
        for t in d.teams:
            name = t["name"]
            ss = d.team_summary(name)
            self.team_combo.addItem(
                f"{name}  •  {d.meta.get('currency','৳')}{money(ss['remaining'])} left  •  {ss['players']}/{d.meta['max_players_per_team']}",
                name,
            )
        target = old_team or d.meta.get("my_team", "")
        idx = self.team_combo.findData(target)
        if idx >= 0: self.team_combo.setCurrentIndex(idx)
        self.team_combo.blockSignals(False)

        self.refresh_player_list()
        self.update_budget_info()
        self.embedded_teams.refresh()

    def matching_players(self):
        d = self.app.model
        q = self.search.text().strip().lower()
        role = self.role_filter.currentText()
        cat = self.category_filter.currentText()
        rows = []
        for idx, p in enumerate(d.players, start=1):
            if d.player_state(p["id"])["status"] != "AVAILABLE": continue
            if role != "All Roles" and p.get("position", "") != role: continue
            if cat != "All Categories" and p.get("category", "") != cat: continue
            if q and q not in f"{idx} {p.get('name','')} {p.get('position','')} {p.get('category','')}".lower(): continue
            rows.append((idx, p))
        rows.sort(key=lambda x: (not x[1].get("wishlist", False), x[0]))
        return rows

    def refresh_player_list(self):
        d = self.app.model
        rows = self.matching_players()
        self.available_count.setText(f"{len(d.available_players())} available")
        self.player_table.blockSignals(True)
        self.player_table.setRowCount(len(rows))
        selected = -1
        cur = d.meta.get("currency", "৳")
        for r, (serial, p) in enumerate(rows):
            maxbid = float(p.get("max_bid", 0) or 0)
            vals = [
                str(serial), p["name"], p.get("position", "") or "—",
                f"{cur}{money(p.get('base_price',0))}",
                f"{cur}{money(maxbid)}" if maxbid else "—",
                "★" if p.get("wishlist") else "",
            ]
            for c, v in enumerate(vals):
                item = table_item(v)
                item.setData(Qt.UserRole, p["id"])
                self.player_table.setItem(r, c, item)
            if p["id"] == self.selected_pid: selected = r
        self.player_table.blockSignals(False)
        if selected < 0 and rows:
            selected = 0
            self.selected_pid = rows[0][1]["id"]
        if selected >= 0: self.player_table.selectRow(selected)
        else: self.selected_pid = None
        self.update_player_card(reset_price=True)

    def player_selected(self):
        r = self.player_table.currentRow()
        item = self.player_table.item(r, 1) if r >= 0 else None
        if item:
            self.selected_pid = item.data(Qt.UserRole)
            self.update_player_card(reset_price=True)

    def update_player_card(self, reset_price=False):
        d = self.app.model
        p = d.get_player(self.selected_pid) if self.selected_pid else None
        cur = d.meta.get("currency", "৳")
        if not p:
            self.player_name.setText("No available player")
            self.player_meta.setText("Search or filter to select a player")
            self.base_price.setText("—")
            self.personal_max.setText("—")
            if reset_price: self.price.setValue(0)
            return
        self.player_name.setText(p["name"])
        meta = " • ".join([x for x in [p.get("position", ""), p.get("category", ""), "Target" if p.get("wishlist") else "Available"] if x])
        self.player_meta.setText(meta)
        self.base_price.setText(f"{cur}{money(p.get('base_price',0))}")
        mb = float(p.get("max_bid", 0) or 0)
        self.personal_max.setText(f"{cur}{money(mb)}" if mb else "Not set")
        if reset_price: self.price.setValue(float(p.get("base_price", 0)))
        self.price_changed()

    def price_changed(self):
        p = self.app.model.get_player(self.selected_pid) if self.selected_pid else None
        mb = float(p.get("max_bid", 0) or 0) if p else 0
        val = self.price.value()
        cur = self.app.model.meta.get("currency", "৳")
        if mb and val > mb:
            self.limit_warning.setText(f"Above your personal max by {cur}{money(val-mb)}")
            self.limit_warning.setStyleSheet("color:#dc2626; font-weight:700;")
        elif mb and val >= mb * 0.9:
            self.limit_warning.setText(f"Near personal max ({cur}{money(mb)})")
            self.limit_warning.setStyleSheet("color:#d97706; font-weight:700;")
        else:
            self.limit_warning.setText("")
            self.limit_warning.setStyleSheet("")

    def update_budget_info(self):
        team = self.current_team()
        if not team:
            self.budget_info.setText("")
            self.safe_bid.setText("—")
            return
        d = self.app.model
        cur = d.meta.get("currency", "৳")
        sm = d.team_summary(team)
        self.budget_info.setText(
            f"Remaining {cur}{money(sm['remaining'])}   •   Players {sm['players']}/{d.meta['max_players_per_team']}   •   Slots {sm['slots']}"
        )
        self.safe_bid.setText(f"{cur}{money(d.safe_max_bid(team))}")

    def mark_sold(self):
        if not self.selected_pid:
            QMessageBox.information(self, "No Player", "Select an available player first.")
            return
        try:
            self.app.model.record_sale(self.selected_pid, self.current_team(), self.price.value())
            self.selected_pid = None
            self.app.after_change(auto_save=True)
        except Exception as e:
            QMessageBox.warning(self, "Cannot Record Sale", str(e))

    def mark_unsold(self):
        if not self.selected_pid:
            QMessageBox.information(self, "No Player", "Select an available player first.")
            return
        try:
            self.app.model.record_unsold(self.selected_pid)
            self.selected_pid = None
            self.app.after_change(auto_save=True)
        except Exception as e:
            QMessageBox.warning(self, "Cannot Mark Unsold", str(e))

    def undo_last(self):
        if not self.app.model.undo_last():
            QMessageBox.information(self, "Undo", "There is nothing to undo.")
            return
        self.app.after_change(auto_save=True)

class PlayersPage(QWidget):
    def __init__(self, app):
        super().__init__(); self.app=app
        root=QVBoxLayout(self); root.setContentsMargins(24,18,24,18); root.setSpacing(10)
        header=QHBoxLayout(); title=QLabel("Players"); title.setObjectName("SectionTitle"); header.addWidget(title); header.addStretch()
        self.search=QLineEdit(); self.search.setPlaceholderText("Search name, role, category or team..."); self.search.setClearButtonEnabled(True); self.search.setMaximumWidth(300); self.search.textChanged.connect(self.refresh); header.addWidget(self.search)
        self.status_filter=QComboBox(); self.status_filter.addItems(["All","Available","Sold","Unsold","My Players","Wishlist"]); self.status_filter.currentTextChanged.connect(self.refresh); header.addWidget(self.status_filter)
        baseall=QPushButton("Set Base Price for All"); baseall.clicked.connect(self.set_base_all); header.addWidget(baseall)
        add=QPushButton("Add Player"); add.setIcon(self.app.icon("plus")); add.setObjectName("Primary"); add.clicked.connect(self.add_player); header.addWidget(add); root.addLayout(header)

        filt=QHBoxLayout(); self.role_filter=QComboBox(); self.category_filter=QComboBox(); self.role_filter.currentTextChanged.connect(self.refresh); self.category_filter.currentTextChanged.connect(self.refresh); filt.addWidget(QLabel("Role")); filt.addWidget(self.role_filter); filt.addWidget(QLabel("Category")); filt.addWidget(self.category_filter); filt.addStretch(); self.count=QLabel(); self.count.setObjectName("Muted"); filt.addWidget(self.count); root.addLayout(filt)

        bulk=QHBoxLayout(); bulk.addWidget(QLabel("Bulk actions:")); wish=QPushButton("Add Wishlist"); wish.setIcon(self.app.icon("star")); wish.clicked.connect(lambda:self.bulk_wishlist(True)); unwish=QPushButton("Remove Wishlist"); unwish.clicked.connect(lambda:self.bulk_wishlist(False)); setbase=QPushButton("Set Base Price"); setbase.clicked.connect(self.bulk_base); delete=QPushButton("Delete Selected"); delete.setIcon(self.app.icon("trash")); delete.clicked.connect(self.bulk_delete)
        for b in (wish,unwish,setbase,delete):bulk.addWidget(b)
        bulk.addStretch(); root.addLayout(bulk)

        self.table=QTableWidget(0,10); self.table.setHorizontalHeaderLabels(["Player","Position","Category","Base Price","Status","Team","Sold Price","My Max","Target","Actions"])
        self.table.horizontalHeader().setSectionResizeMode(0,QHeaderView.Stretch); self.table.horizontalHeader().setSectionResizeMode(1,QHeaderView.Stretch); self.table.horizontalHeader().setSectionResizeMode(5,QHeaderView.Stretch)
        for c in (2,3,4,6,7,8): self.table.horizontalHeader().setSectionResizeMode(c,QHeaderView.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(9,QHeaderView.Fixed); self.table.setColumnWidth(9,142)
        self.table.verticalHeader().setVisible(False); self.table.setAlternatingRowColors(True); self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setSelectionMode(QAbstractItemView.ExtendedSelection); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers); self.table.doubleClicked.connect(self.edit_selected); root.addWidget(self.table,1)

    def sync_filters(self):
        d=self.app.model; old_role,old_cat=self.role_filter.currentText(),self.category_filter.currentText(); roles=sorted({p.get("position","").strip() for p in d.players if p.get("position","").strip()}); cats=sorted({p.get("category","").strip() for p in d.players if p.get("category","").strip()})
        self.role_filter.blockSignals(True); self.category_filter.blockSignals(True); self.role_filter.clear(); self.role_filter.addItem("All Roles"); self.role_filter.addItems(roles); self.category_filter.clear(); self.category_filter.addItem("All Categories"); self.category_filter.addItems(cats)
        if old_role in roles:self.role_filter.setCurrentText(old_role)
        if old_cat in cats:self.category_filter.setCurrentText(old_cat)
        self.role_filter.blockSignals(False); self.category_filter.blockSignals(False)

    def filtered_players(self):
        d=self.app.model; q=self.search.text().strip().lower(); filt=self.status_filter.currentText(); role=self.role_filter.currentText(); cat=self.category_filter.currentText(); out=[]
        for p in d.players:
            st=d.player_state(p["id"]); hay=" ".join([p.get("name",""),p.get("position",""),p.get("category",""),st.get("team","")]).lower()
            if q and q not in hay:continue
            if role not in ("","All Roles") and p.get("position","")!=role:continue
            if cat not in ("","All Categories") and p.get("category","")!=cat:continue
            if filt=="Available" and st["status"]!="AVAILABLE":continue
            if filt=="Sold" and st["status"]!="SOLD":continue
            if filt=="Unsold" and st["status"]!="UNSOLD":continue
            if filt=="My Players" and st.get("team")!=d.meta.get("my_team"):continue
            if filt=="Wishlist" and not p.get("wishlist",False):continue
            out.append((p,st))
        return out

    def refresh(self):
        d=self.app.model; self.sync_filters(); cur=d.meta.get("currency","৳"); rows=self.filtered_players(); self.table.setRowCount(len(rows))
        for r,(p,st) in enumerate(rows):
            mb=float(p.get("max_bid",0) or 0); vals=[p["name"],p.get("position","") or "—",p.get("category","") or "—",f"{cur}{money(p.get('base_price',0))}",st["status"],st.get("team","") or "—",f"{cur}{money(st.get('price',0))}" if st["status"]=="SOLD" else "—",f"{cur}{money(mb)}" if mb else "—","Yes" if p.get("wishlist") else "—"]
            for c,v in enumerate(vals): item=table_item(v); item.setData(Qt.UserRole,p["id"]); self.table.setItem(r,c,item)
            star=icon_button(self.app,"star","Toggle wishlist",lambda checked=False,pid=p["id"]:self.toggle_wishlist(pid)); edit=icon_button(self.app,"edit","Edit player",lambda checked=False,pid=p["id"]:self.edit_player(pid)); delete=icon_button(self.app,"trash","Delete player",lambda checked=False,pid=p["id"]:self.delete_player(pid),danger=True); self.table.setCellWidget(r,9,action_cell(star,edit,delete))
        self.count.setText(f"{len(rows)} shown • {len(d.players)} total")

    def selected_pids(self):
        pids=[]
        for idx in self.table.selectionModel().selectedRows(0):
            item=self.table.item(idx.row(),0)
            if item and item.data(Qt.UserRole) not in pids:pids.append(item.data(Qt.UserRole))
        return pids
    def selected_pid(self):
        pids=self.selected_pids(); return pids[0] if pids else None
    def add_player(self):
        d=self.app.model; dlg=PlayerDialog(self,common_base_enabled=d.meta.get("use_common_base_price",False),common_base_price=d.meta.get("common_base_price",0))
        if dlg.exec():d.add_player(**dlg.values()); self.app.after_change(auto_save=True)
    def edit_selected(self):
        pid=self.selected_pid();
        if pid:self.edit_player(pid)
    def edit_player(self,pid):
        d=self.app.model; p=d.get_player(pid)
        if not p:return
        dlg=PlayerDialog(self,p,d.meta.get("use_common_base_price",False),d.meta.get("common_base_price",0))
        if dlg.exec():d.update_player(pid,**dlg.values()); self.app.after_change(auto_save=True)
    def toggle_wishlist(self,pid):
        p=self.app.model.get_player(pid)
        if p:self.app.model.set_wishlist(pid,not p.get("wishlist",False)); self.app.after_change(auto_save=True)
    def delete_player(self,pid):
        p=self.app.model.get_player(pid)
        if not p:return
        if QMessageBox.question(self,"Delete Player",f"Delete {p['name']}?")!=QMessageBox.Yes:return
        try:self.app.model.delete_player(pid); self.app.after_change(auto_save=True)
        except Exception as e:QMessageBox.warning(self,"Cannot Delete",str(e))
    def set_base_all(self):
        val,ok=QInputDialog.getDouble(self,"Set Base Price for All","Base price",float(self.app.model.meta.get("common_base_price",0) or 0),0,1_000_000_000_000,0)
        if ok:self.app.model.set_all_base_prices(val); self.app.after_change(auto_save=True)
    def bulk_wishlist(self,enabled):
        for pid in self.selected_pids():self.app.model.set_wishlist(pid,enabled)
        self.app.after_change(auto_save=True)
    def bulk_base(self):
        pids=self.selected_pids()
        if not pids:return
        val,ok=QInputDialog.getDouble(self,"Set Base Price","Base price for selected players",0,0,1_000_000_000_000,0)
        if ok:
            for pid in pids:self.app.model.update_player(pid,base_price=val)
            self.app.after_change(auto_save=True)
    def bulk_delete(self):
        pids=self.selected_pids()
        if not pids:return
        if QMessageBox.question(self,"Delete Players",f"Delete {len(pids)} selected player(s)? Players with auction history will be kept.")!=QMessageBox.Yes:return
        errors=0
        for pid in pids:
            try:self.app.model.delete_player(pid)
            except Exception:errors+=1
        self.app.after_change(auto_save=True)
        if errors:QMessageBox.information(self,"Bulk Delete",f"{errors} player(s) could not be deleted because they have auction history.")


class TeamsPage(QWidget):
    """Team table + selected-team detail. Also embedded under Live Tracker."""

    def __init__(self, app, embedded=False):
        super().__init__()
        self.app = app
        self.embedded = embedded

        root = QVBoxLayout(self)
        root.setContentsMargins(4 if embedded else 24, 5 if embedded else 18, 4 if embedded else 24, 5 if embedded else 18)
        root.setSpacing(7 if embedded else 10)

        head = QHBoxLayout()
        icon = QLabel()
        icon.setPixmap(self.app.icon("teams").pixmap(22, 22))
        head.addWidget(icon)
        title = QLabel("Teams")
        title.setObjectName("SectionTitle")
        head.addWidget(title)
        if embedded:
            drag_hint = QLabel("Drag the divider above to resize this section")
            drag_hint.setObjectName("CompactHint")
            head.addWidget(drag_hint)
        head.addStretch()
        add = QPushButton("Add Team")
        add.setIcon(self.app.icon("plus"))
        add.setObjectName("Primary")
        add.clicked.connect(self.add_team)
        head.addWidget(add)
        root.addLayout(head)

        self.split = QSplitter(Qt.Horizontal)
        self.split.setChildrenCollapsible(False)
        self.split.setHandleWidth(7)

        self.table = QTableWidget(0, 7)
        self.table.setHorizontalHeaderLabels(["Team", "Players", "Spent", "Remaining", "Slots Left", "Safe Max Bid", "Actions"])
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        for c in range(1, 6): self.table.horizontalHeader().setSectionResizeMode(c, QHeaderView.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(6, QHeaderView.Fixed)
        self.table.setColumnWidth(6, 132)
        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setDefaultSectionSize(42)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.doubleClicked.connect(self.edit_selected_team)
        self.table.itemSelectionChanged.connect(self.update_detail)
        self.split.addWidget(self.table)

        detail = QFrame()
        detail.setObjectName("Card")
        detail.setMinimumWidth(285)
        dl = QVBoxLayout(detail)
        dl.setContentsMargins(12, 10, 12, 10)
        dl.setSpacing(6)
        self.detail_title = QLabel("Select a team")
        self.detail_title.setObjectName("SectionTitle")
        dl.addWidget(self.detail_title)
        self.detail_metrics = QLabel()
        self.detail_metrics.setWordWrap(True)
        self.detail_metrics.setObjectName("Muted")
        dl.addWidget(self.detail_metrics)
        squad_title = QLabel("Squad")
        squad_title.setObjectName("Brand")
        dl.addWidget(squad_title)
        self.squad = QTableWidget(0, 3)
        self.squad.setHorizontalHeaderLabels(["Player", "Role", "Price"])
        self.squad.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.squad.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.squad.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeToContents)
        self.squad.verticalHeader().setVisible(False)
        self.squad.verticalHeader().setDefaultSectionSize(38)
        self.squad.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.squad.setAlternatingRowColors(True)
        dl.addWidget(self.squad, 1)
        self.split.addWidget(detail)
        self.split.setStretchFactor(0, 3)
        self.split.setStretchFactor(1, 1)
        self.split.setSizes([900, 360])
        root.addWidget(self.split, 1)

    def refresh(self):
        d = self.app.model
        cur = d.meta.get("currency", "৳")
        selected = self.selected_team()
        self.table.setRowCount(len(d.teams))
        for r, t in enumerate(d.teams):
            name = t["name"]
            sm = d.team_summary(name)
            vals = [
                f"{name}  (My Team)" if name == d.meta.get("my_team") else name,
                f"{sm['players']}/{d.meta['max_players_per_team']}",
                f"{cur}{money(sm['spent'])}",
                f"{cur}{money(sm['remaining'])}",
                str(sm["slots"]),
                f"{cur}{money(d.safe_max_bid(name))}",
            ]
            for c, v in enumerate(vals):
                item = table_item(v)
                item.setData(Qt.UserRole, name)
                if name == d.meta.get("my_team"):
                    f = item.font(); f.setBold(True); item.setFont(f)
                self.table.setItem(r, c, item)
            edit = icon_button(self.app, "edit", "Rename team", lambda checked=False, n=name: self.edit_team(n))
            my = icon_button(self.app, "myteam", "Set as My Team", lambda checked=False, n=name: self.set_my_team(n))
            delete = icon_button(self.app, "trash", "Delete team", lambda checked=False, n=name: self.delete_team(n), danger=True)
            self.table.setCellWidget(r, 6, action_cell(edit, my, delete))
            if name == selected: self.table.selectRow(r)
        if self.table.currentRow() < 0 and d.teams: self.table.selectRow(0)
        self.update_detail()

    def selected_team(self):
        r = self.table.currentRow()
        item = self.table.item(r, 0) if r >= 0 else None
        return item.data(Qt.UserRole) if item else None

    def update_detail(self):
        name = self.selected_team()
        if not name:
            self.detail_title.setText("Select a team")
            self.detail_metrics.setText("")
            self.squad.setRowCount(0)
            return
        d = self.app.model
        cur = d.meta.get("currency", "৳")
        sm = d.team_summary(name)
        self.detail_title.setText(name + ("  •  My Team" if name == d.meta.get("my_team") else ""))
        self.detail_metrics.setText(
            f"Budget {cur}{money(d.meta.get('budget_per_team',0))}   •   Spent {cur}{money(sm['spent'])}\n"
            f"Remaining {cur}{money(sm['remaining'])}   •   Players {sm['players']}/{d.meta['max_players_per_team']}   •   Safe Bid {cur}{money(d.safe_max_bid(name))}"
        )
        buys = d.team_purchases(name)
        self.squad.setRowCount(len(buys))
        for r, e in enumerate(buys):
            p = d.get_player(e.get("player_id"))
            vals = [p["name"] if p else "Unknown", p.get("position", "") if p else "", f"{cur}{money(e.get('price',0))}"]
            for c, v in enumerate(vals): self.squad.setItem(r, c, table_item(v))

    def add_team(self):
        name, ok = QInputDialog.getText(self, "Add Team", "Team name")
        if ok and name.strip():
            try:
                self.app.model.add_team(name)
                self.app.after_change(auto_save=True)
            except Exception as e: QMessageBox.warning(self, "Cannot Add Team", str(e))

    def edit_selected_team(self):
        name = self.selected_team()
        if name: self.edit_team(name)

    def edit_team(self, old):
        new, ok = QInputDialog.getText(self, "Edit Team", "Team name", text=old)
        if ok:
            try:
                self.app.model.rename_team(old, new)
                self.app.after_change(auto_save=True)
            except Exception as e: QMessageBox.warning(self, "Cannot Edit Team", str(e))

    def set_my_team(self, name):
        self.app.model.meta["my_team"] = name
        self.app.after_change(auto_save=True)

    def delete_team(self, name):
        sm = self.app.model.team_summary(name)
        cur = self.app.model.meta.get("currency", "৳")
        if QMessageBox.question(self, "Delete Team", f"Delete {name}?\nPlayers: {sm['players']}\nSpent: {cur}{money(sm['spent'])}") != QMessageBox.Yes:
            return
        try:
            self.app.model.delete_team(name)
            self.app.after_change(auto_save=True)
        except Exception as e: QMessageBox.warning(self, "Cannot Delete Team", str(e))

class MyTeamPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 22, 24, 22)
        root.setSpacing(15)

        self.title = QLabel("My Team")
        self.title.setObjectName("SectionTitle")
        root.addWidget(self.title)

        cards = QGridLayout()
        self.budget = MetricCard("Starting Budget")
        self.spent = MetricCard("Spent")
        self.remaining = MetricCard("Remaining")
        self.players = MetricCard("Players")
        self.slots = MetricCard("Slots Left")
        self.safe = MetricCard("Safe Max Bid")
        for i, c in enumerate([self.budget, self.spent, self.remaining, self.players, self.slots, self.safe]):
            cards.addWidget(c, i // 3, i % 3)
        root.addLayout(cards)

        self.progress = QProgressBar()
        self.progress.setTextVisible(False)
        root.addWidget(self.progress)

        label = QLabel("Squad")
        label.setObjectName("SectionTitle")
        root.addWidget(label)

        self.table = QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels(["Player", "Position", "Category", "Base Price", "Bought For"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        root.addWidget(self.table, 1)

    def refresh(self):
        d = self.app.model
        team = d.meta.get("my_team", "")
        cur = d.meta.get("currency", "৳")
        s = d.team_summary(team) if team else {"players":0, "spent":0, "remaining":0, "slots":0}
        self.title.setText(f"My Team — {team or 'Not Selected'}")
        self.budget.value.setText(f"{cur}{money(d.meta.get('budget_per_team',0))}")
        self.spent.value.setText(f"{cur}{money(s['spent'])}")
        self.remaining.value.setText(f"{cur}{money(s['remaining'])}")
        self.players.value.setText(f"{s['players']}/{d.meta.get('max_players_per_team',0)}")
        self.slots.value.setText(str(s["slots"]))
        self.safe.value.setText(f"{cur}{money(d.safe_max_bid(team) if team else 0)}")

        budget = float(d.meta.get("budget_per_team", 0))
        pct = int((s["spent"] / budget) * 100) if budget else 0
        self.progress.setValue(max(0, min(100, pct)))

        purchases = [e for e in d.events if e.get("type") == "sale" and e.get("team") == team]
        self.table.setRowCount(len(purchases))
        for r, e in enumerate(purchases):
            p = d.get_player(e.get("player_id"))
            if not p:
                continue
            vals = [p["name"], p.get("position",""), p.get("category",""),
                    f"{cur}{money(p.get('base_price',0))}", f"{cur}{money(e.get('price',0))}"]
            for c, v in enumerate(vals):
                self.table.setItem(r, c, table_item(v))


class WishlistPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 22, 24, 22)
        root.setSpacing(12)
        head=QHBoxLayout()
        title=QLabel("Wishlist / Targets")
        title.setObjectName("SectionTitle")
        head.addWidget(title)
        head.addStretch()
        self.search=QLineEdit()
        self.search.setPlaceholderText("Search target players...")
        self.search.setClearButtonEnabled(True)
        self.search.textChanged.connect(self.refresh)
        self.search.setMaximumWidth(300)
        head.addWidget(self.search)
        root.addLayout(head)

        hint=QLabel("Target players appear first in the Live Tracker player list. Add/remove targets from the Players page.")
        hint.setObjectName("Muted")
        root.addWidget(hint)

        self.table=QTableWidget(0,7)
        self.table.setHorizontalHeaderLabels(["Player","Position","Category","Base Price","My Max Bid","Status","Actions"])
        for c in range(6): self.table.horizontalHeader().setSectionResizeMode(c,QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(6,QHeaderView.Fixed); self.table.setColumnWidth(6,96)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        root.addWidget(self.table,1)

    def refresh(self):
        d=self.app.model
        cur=d.meta.get("currency","৳")
        q=self.search.text().strip().lower()
        rows=[]
        for p in d.wishlist_players():
            st=d.player_state(p["id"])
            hay=f"{p.get('name','')} {p.get('position','')} {p.get('category','')}".lower()
            if q and q not in hay: continue
            rows.append((p,st))
        self.table.setRowCount(len(rows))
        for r,(p,st) in enumerate(rows):
            mb=float(p.get("max_bid",0) or 0)
            vals=[p["name"],p.get("position","") or "—",p.get("category","") or "—",f"{cur}{money(p.get('base_price',0))}",f"{cur}{money(mb)}" if mb else "—",st["status"]]
            for c,v in enumerate(vals): self.table.setItem(r,c,table_item(v))
            edit=icon_button(self.app,"edit","Edit target / personal max bid",lambda checked=False,pid=p["id"]: self.edit_target(pid))
            remove=icon_button(self.app,"trash","Remove from wishlist",lambda checked=False,pid=p["id"]: self.remove(pid))
            self.table.setCellWidget(r,6,action_cell(edit,remove))

    def edit_target(self,pid):
        d=self.app.model; p=d.get_player(pid)
        if not p:return
        dlg=PlayerDialog(self,p,d.meta.get("use_common_base_price",False),d.meta.get("common_base_price",0))
        if dlg.exec():d.update_player(pid,**dlg.values()); self.app.after_change(auto_save=True)

    def remove(self,pid):
        self.app.model.set_wishlist(pid,False)
        self.app.after_change(auto_save=True)


class HistoryPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 22, 24, 22)
        root.setSpacing(12)

        head = QHBoxLayout()
        title = QLabel("History")
        title.setObjectName("SectionTitle")
        head.addWidget(title)
        head.addStretch()
        self.search=QLineEdit()
        self.search.setPlaceholderText("Search player or team...")
        self.search.setClearButtonEnabled(True)
        self.search.textChanged.connect(self.refresh)
        self.search.setMaximumWidth(280)
        head.addWidget(self.search)
        root.addLayout(head)

        self.table = QTableWidget(0, 6)
        self.table.setHorizontalHeaderLabels(["Time", "Player", "Status", "Team", "Price", "Actions"])
        for c in range(5): self.table.horizontalHeader().setSectionResizeMode(c, QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(5, QHeaderView.Fixed); self.table.setColumnWidth(5, 92)
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.doubleClicked.connect(self.edit_selected)
        root.addWidget(self.table, 1)

    def refresh(self):
        d = self.app.model
        cur = d.meta.get("currency", "৳")
        q=self.search.text().strip().lower()
        events=[]
        for e in reversed(d.events):
            p=d.get_player(e.get("player_id"))
            pname=p["name"] if p else "Unknown"
            hay=f"{pname} {e.get('team','')} {e.get('type','')}".lower()
            if q and q not in hay: continue
            events.append(e)
        self.table.setRowCount(len(events))
        for r, e in enumerate(events):
            p = d.get_player(e.get("player_id"))
            pname = p["name"] if p else "Unknown"
            vals = [e.get("time","").replace("T"," "), pname, "SOLD" if e["type"] == "sale" else "UNSOLD", e.get("team","") if e["type"] == "sale" else "—", f"{cur}{money(e.get('price',0))}" if e["type"] == "sale" else "—"]
            for c,v in enumerate(vals):
                item=table_item(v); item.setData(Qt.UserRole,e["id"]); self.table.setItem(r,c,item)
            buttons=[]
            if e["type"]=="sale":
                buttons.append(icon_button(self.app,"edit","Edit sale",lambda checked=False,eid=e["id"]: self.edit_event(eid)))
            buttons.append(icon_button(self.app,"undo","Undo this entry",lambda checked=False,eid=e["id"]: self.undo_event(eid)))
            self.table.setCellWidget(r,5,action_cell(*buttons))

    def selected_event(self):
        r=self.table.currentRow()
        if r<0:return None
        eid=self.table.item(r,0).data(Qt.UserRole)
        return next((e for e in self.app.model.events if e["id"]==eid),None)

    def edit_selected(self):
        e=self.selected_event()
        if e: self.edit_event(e["id"])

    def edit_event(self,eid):
        e=next((x for x in self.app.model.events if x["id"]==eid),None)
        if not e:return
        if e["type"]!="sale":
            QMessageBox.information(self,"Edit Entry","Only sold entries have editable team/price values.")
            return
        d=self.app.model
        p=d.get_player(e["player_id"])
        dlg=SaleEditDialog([t["name"] for t in d.teams],p["name"] if p else "Player",e,d.meta.get("currency","৳"),self)
        if dlg.exec():
            original_index=next(i for i,item in enumerate(d.events) if item["id"]==eid)
            original=dict(e); edited=dict(e); edited["team"]=dlg.team.currentText(); edited["price"]=dlg.price.value()
            d.events.pop(original_index)
            try:
                if edited["price"]>d.team_summary(edited["team"])["remaining"]:
                    raise ValueError(f"{edited['team']} has only {d.meta['currency']}{money(d.team_summary(edited['team'])['remaining'])} remaining.")
                if d.team_summary(edited["team"])["players"]>=d.meta["max_players_per_team"]:
                    raise ValueError(f"{edited['team']} has no player slots remaining.")
                d.events.insert(original_index,edited)
            except Exception as ex:
                d.events.insert(original_index,original)
                QMessageBox.warning(self,"Invalid Edit",str(ex)); self.app.refresh_all(); return
            self.app.after_change(auto_save=True)

    def undo_event(self,eid):
        e=next((x for x in self.app.model.events if x["id"]==eid),None)
        if not e:return
        p=self.app.model.get_player(e.get("player_id")); name=p["name"] if p else "this player"
        if QMessageBox.question(self,"Undo Entry",f"Undo auction entry for {name}?\nThe player will become available again.")!=QMessageBox.Yes:return
        self.app.model.delete_event(eid)
        self.app.after_change(auto_save=True)

    def remove_selected(self):
        e=self.selected_event()
        if e:self.undo_event(e["id"])


class ExcelColumnMappingDialog(QDialog):
    FIELD_SPECS = [
        ("name", "Player Name *", ["name", "player", "player name", "player_name", "full name"]),
        ("position", "Position / Role", ["position", "role", "playing role", "post"]),
        ("category", "Category", ["category", "grade", "type", "class"]),
        ("base_price", "Base Price", ["base price", "base_price", "starting price", "starting value", "price", "base"]),
        ("max_bid", "Personal Max Bid", ["max bid", "max_bid", "personal max bid", "my max", "maximum bid"]),
        ("wishlist", "Wishlist / Target", ["wishlist", "target", "favourite", "favorite", "shortlist"]),
        ("notes", "Notes", ["notes", "note", "remarks", "remark", "comment"]),
    ]

    def __init__(self, path, parent=None):
        super().__init__(parent)
        self.path = path
        self.workbook = load_workbook(path, data_only=True)
        self.setWindowTitle("Map Excel Columns")
        self.resize(880, 690)

        root = QVBoxLayout(self)
        root.setSpacing(12)

        title = QLabel("Map Excel Columns")
        title.setObjectName("SectionTitle")
        root.addWidget(title)

        help_text = QLabel(
            "Only Player Name is required. Map whatever columns your file contains; leave missing fields as Not mapped. "
            "Unmapped Excel columns are ignored."
        )
        help_text.setObjectName("Muted")
        help_text.setWordWrap(True)
        root.addWidget(help_text)

        top = QHBoxLayout()
        top.addWidget(QLabel("Sheet"))
        self.sheet_combo = QComboBox()
        self.sheet_combo.addItems(self.workbook.sheetnames)
        self.sheet_combo.currentTextChanged.connect(self.reload_sheet)
        top.addWidget(self.sheet_combo, 1)
        root.addLayout(top)

        mapping_box = QGroupBox("Column Mapping")
        mapping_form = QFormLayout(mapping_box)
        mapping_form.setFieldGrowthPolicy(QFormLayout.AllNonFixedFieldsGrow)
        self.mapping_boxes = {}
        for key, label, _aliases in self.FIELD_SPECS:
            combo = QComboBox()
            self.mapping_boxes[key] = combo
            mapping_form.addRow(label, combo)
        root.addWidget(mapping_box)

        preview_label = QLabel("Excel Preview")
        preview_label.setObjectName("SectionTitle")
        root.addWidget(preview_label)

        self.preview = QTableWidget(0, 0)
        self.preview.verticalHeader().setVisible(False)
        self.preview.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.preview.setAlternatingRowColors(True)
        self.preview.setMaximumHeight(260)
        root.addWidget(self.preview, 1)

        note = QLabel(
            "If Base Price is not mapped, the app uses your auction common base price when enabled; otherwise it imports 0. "
            "The same applies to blank/invalid base-price cells."
        )
        note.setObjectName("Muted")
        note.setWordWrap(True)
        root.addWidget(note)

        buttons = QDialogButtonBox(QDialogButtonBox.Cancel | QDialogButtonBox.Ok)
        buttons.button(QDialogButtonBox.Ok).setText("Continue to Preview")
        buttons.accepted.connect(self.validate_and_accept)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

        self.reload_sheet()

    @staticmethod
    def _normalize(value):
        return " ".join(str(value or "").strip().lower().replace("_", " ").replace("-", " ").split())

    def current_headers(self):
        ws = self.workbook[self.sheet_combo.currentText()]
        headers = []
        for c in range(1, ws.max_column + 1):
            raw = ws.cell(1, c).value
            label = str(raw).strip() if raw is not None and str(raw).strip() else f"Column {c}"
            headers.append((c, label))
        return headers

    def reload_sheet(self):
        headers = self.current_headers()
        header_labels = [label for _, label in headers]

        for key, _label, aliases in self.FIELD_SPECS:
            combo = self.mapping_boxes[key]
            combo.blockSignals(True)
            combo.clear()
            combo.addItem("— Not mapped —", None)
            for col_idx, header in headers:
                combo.addItem(header, col_idx)

            alias_norm = {self._normalize(a) for a in aliases}
            detected = None
            for col_idx, header in headers:
                if self._normalize(header) in alias_norm:
                    detected = col_idx
                    break
            if detected is not None:
                idx = combo.findData(detected)
                if idx >= 0:
                    combo.setCurrentIndex(idx)
            combo.blockSignals(False)

        ws = self.workbook[self.sheet_combo.currentText()]
        sample_rows = min(max(ws.max_row - 1, 0), 6)
        self.preview.clear()
        self.preview.setColumnCount(len(headers))
        self.preview.setHorizontalHeaderLabels(header_labels)
        self.preview.setRowCount(sample_rows)
        for r in range(sample_rows):
            excel_row = r + 2
            for c, (col_idx, _header) in enumerate(headers):
                value = ws.cell(excel_row, col_idx).value
                self.preview.setItem(r, c, table_item("" if value is None else value))
        if headers:
            self.preview.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
            self.preview.horizontalHeader().setStretchLastSection(True)

    def validate_and_accept(self):
        if self.mapping_boxes["name"].currentData() is None:
            QMessageBox.warning(self, "Player Name Required", "Map one Excel column to Player Name before continuing.")
            return

        used = []
        for key, _label, _aliases in self.FIELD_SPECS:
            data = self.mapping_boxes[key].currentData()
            if data is not None:
                used.append((key, data))
        cols = [c for _, c in used]
        if len(cols) != len(set(cols)):
            QMessageBox.warning(self, "Duplicate Mapping", "Each Excel column can map to only one app field.")
            return
        self.accept()

    def result_mapping(self):
        return {
            "sheet": self.sheet_combo.currentText(),
            "fields": {key: combo.currentData() for key, combo in self.mapping_boxes.items()},
            "labels": {key: combo.currentText() for key, combo in self.mapping_boxes.items()},
        }


class ImportPreviewDialog(QDialog):
    def __init__(self, rows, currency, mapping_summary="", ignored_blank=0, parent=None):
        super().__init__(parent)
        self.rows = rows
        self.setWindowTitle("Import Preview")
        self.resize(950, 650)
        root = QVBoxLayout(self)
        root.setSpacing(10)

        title = QLabel("Import Preview")
        title.setObjectName("SectionTitle")
        root.addWidget(title)

        valid = sum(1 for x in rows if x["status"] == "VALID")
        dup = sum(1 for x in rows if x["status"] == "DUPLICATE")
        defaults = sum(1 for x in rows if x.get("import_note"))
        summary = QLabel(
            f"Ready: {valid}   •   Duplicates: {dup}   •   Defaults/Warnings: {defaults}   •   Blank-name rows ignored: {ignored_blank}"
        )
        summary.setObjectName("Muted")
        root.addWidget(summary)

        if mapping_summary:
            mapped = QLabel(mapping_summary)
            mapped.setObjectName("Muted")
            mapped.setWordWrap(True)
            root.addWidget(mapped)

        self.table = QTableWidget(0, 8)
        self.table.setHorizontalHeaderLabels(["Row", "Name", "Position", "Category", "Base Price", "My Max", "Status", "Import Note"])
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(3, QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(7, QHeaderView.Stretch)
        for c in (0, 4, 5, 6):
            self.table.horizontalHeader().setSectionResizeMode(c, QHeaderView.ResizeToContents)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        preview = rows[:300]
        self.table.setRowCount(len(preview))
        for r, x in enumerate(preview):
            vals = [
                x["row"],
                x.get("name", ""),
                x.get("position", "") or "—",
                x.get("category", "") or "—",
                f"{currency}{money(x.get('base_price', 0))}",
                f"{currency}{money(x.get('max_bid', 0))}" if x.get("max_bid", 0) else "—",
                x["status"],
                x.get("import_note", "") or "—",
            ]
            for c, v in enumerate(vals):
                self.table.setItem(r, c, table_item(v))
        root.addWidget(self.table, 1)

        opts = QFormLayout()
        self.duplicate_policy = QComboBox()
        self.duplicate_policy.addItems(["Skip duplicates", "Replace existing duplicates", "Import duplicates anyway"])
        self.common = QCheckBox("Override every imported player's base price with one common value")
        self.common_price = QDoubleSpinBox()
        self.common_price.setRange(0, 1_000_000_000_000)
        self.common_price.setDecimals(0)
        self.common_price.setSingleStep(500)
        self.common_price.setEnabled(False)
        self.common.toggled.connect(self.common_price.setEnabled)
        opts.addRow("Duplicate handling", self.duplicate_policy)
        opts.addRow("Common base price", self.common)
        opts.addRow("Common price amount", self.common_price)
        root.addLayout(opts)

        buttons = QDialogButtonBox(QDialogButtonBox.Cancel | QDialogButtonBox.Ok)
        buttons.button(QDialogButtonBox.Ok).setText("Import Players")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)


class ImportExportPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 22, 24, 22)
        root.setSpacing(16)

        title = QLabel("Import / Export")
        title.setObjectName("SectionTitle")
        root.addWidget(title)

        g1 = QGroupBox("Player List")
        l1 = QVBoxLayout(g1)
        txt = QLabel(
            "Import almost any .xlsx player list using column mapping. Only Player Name is required; Position, Category, Base Price, Notes, Wishlist and Personal Max Bid are optional."
        )
        txt.setWordWrap(True)
        txt.setObjectName("Muted")
        l1.addWidget(txt)

        flow = QLabel("1. Choose Excel  →  2. Map columns  →  3. Preview  →  4. Import")
        flow.setObjectName("Muted")
        l1.addWidget(flow)

        row = QHBoxLayout()
        imp = QPushButton("Import Players from Excel")
        imp.setIcon(self.app.icon("import"))
        imp.setObjectName("Primary")
        imp.clicked.connect(self.import_excel)
        template = QPushButton("Create Excel Template")
        template.clicked.connect(self.create_template)
        row.addWidget(imp)
        row.addWidget(template)
        row.addStretch()
        l1.addLayout(row)
        root.addWidget(g1)

        g2 = QGroupBox("Auction Results")
        l2 = QVBoxLayout(g2)
        t = QLabel("Export players, team balances and complete history to one Excel workbook. Missing optional player fields are exported as blank/0 values.")
        t.setObjectName("Muted")
        t.setWordWrap(True)
        l2.addWidget(t)
        exp = QPushButton("Export Results to Excel")
        exp.clicked.connect(self.export_excel)
        l2.addWidget(exp, 0, Qt.AlignLeft)
        root.addWidget(g2)
        root.addStretch()

    def refresh(self):
        pass

    @staticmethod
    def _truthy(value):
        return str(value or "").strip().lower() in ("1", "true", "yes", "y", "target", "wishlist", "star", "favorite", "favourite")

    def parse_excel(self, path, mapping):
        d = self.app.model
        wb = load_workbook(path, data_only=True)
        ws = wb[mapping["sheet"]]
        fields = mapping["fields"]
        existing = {p["name"].strip().lower() for p in d.players}
        rows = []
        ignored_blank = 0

        def cell(row_number, field):
            col = fields.get(field)
            if col is None:
                return None
            return ws.cell(row_number, col).value

        global_common = bool(d.meta.get("use_common_base_price", False))
        global_base = float(d.meta.get("common_base_price", 0) or 0)

        for r in range(2, ws.max_row + 1):
            raw_name = cell(r, "name")
            if raw_name is None or not str(raw_name).strip():
                # Ignore fully blank rows and rows without a player name.
                if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1)):
                    ignored_blank += 1
                continue

            name = str(raw_name).strip()
            position = str(cell(r, "position") or "").strip()
            category = str(cell(r, "category") or "").strip()
            notes = str(cell(r, "notes") or "").strip()
            wishlist = self._truthy(cell(r, "wishlist")) if fields.get("wishlist") is not None else False

            import_notes = []
            raw_base = cell(r, "base_price")
            if fields.get("base_price") is None:
                base = global_base if global_common else 0
                import_notes.append("Base price from common setting" if global_common else "No base-price column → 0")
            elif raw_base in (None, ""):
                base = global_base if global_common else 0
                import_notes.append("Blank base price → common" if global_common else "Blank base price → 0")
            else:
                try:
                    base = float(raw_base)
                except Exception:
                    base = global_base if global_common else 0
                    import_notes.append("Invalid base price → common" if global_common else "Invalid base price → 0")

            raw_max = cell(r, "max_bid")
            if fields.get("max_bid") is None or raw_max in (None, ""):
                max_bid = 0
            else:
                try:
                    max_bid = float(raw_max)
                except Exception:
                    max_bid = 0
                    import_notes.append("Invalid My Max → 0")

            status = "DUPLICATE" if name.lower() in existing else "VALID"
            rows.append({
                "row": r,
                "name": name,
                "position": position,
                "category": category,
                "notes": notes,
                "wishlist": wishlist,
                "base_price": base,
                "max_bid": max_bid,
                "status": status,
                "import_note": "; ".join(import_notes),
            })
        return rows, ignored_blank

    def import_excel(self):
        path, _ = QFileDialog.getOpenFileName(self, "Import Player List", "", "Excel Files (*.xlsx)")
        if not path:
            return
        try:
            mapping_dialog = ExcelColumnMappingDialog(path, self)
            if not mapping_dialog.exec():
                return
            mapping = mapping_dialog.result_mapping()
            rows, ignored_blank = self.parse_excel(path, mapping)
            if not rows:
                QMessageBox.warning(self, "Nothing to Import", "No rows with a mapped Player Name were found in the selected sheet.")
                return

            mapped_parts = []
            for key, label, _aliases in ExcelColumnMappingDialog.FIELD_SPECS:
                source = mapping["labels"].get(key, "— Not mapped —")
                if source != "— Not mapped —":
                    mapped_parts.append(f"{label.replace(' *','')} ← {source}")
            mapping_summary = f"Sheet: {mapping['sheet']}   •   " + "   •   ".join(mapped_parts)

            dlg = ImportPreviewDialog(
                rows,
                self.app.model.meta.get("currency", "৳"),
                mapping_summary,
                ignored_blank,
                self,
            )
            if not dlg.exec():
                return

            d = self.app.model
            policy = dlg.duplicate_policy.currentText()
            common = dlg.common.isChecked()
            common_price = dlg.common_price.value()
            imported = replaced = skipped = 0

            for x in rows:
                existing = next((p for p in d.players if p["name"].strip().lower() == x["name"].lower()), None)
                base = common_price if common else x["base_price"]
                if existing:
                    if policy == "Skip duplicates":
                        skipped += 1
                        continue
                    if policy == "Replace existing duplicates":
                        d.update_player(
                            existing["id"],
                            name=x["name"],
                            position=x["position"],
                            category=x["category"],
                            base_price=base,
                            notes=x["notes"],
                            wishlist=x["wishlist"],
                            max_bid=x["max_bid"],
                        )
                        replaced += 1
                        continue
                d.add_player(x["name"], x["position"], base, x["category"], x["notes"], x["wishlist"], x["max_bid"])
                imported += 1

            self.app.after_change(auto_save=True)
            QMessageBox.information(
                self,
                "Import Complete",
                f"Imported: {imported}\nReplaced: {replaced}\nSkipped duplicates: {skipped}\nBlank-name rows ignored: {ignored_blank}",
            )
        except Exception as e:
            QMessageBox.critical(self, "Import Failed", str(e))

    def create_template(self):
        path, _ = QFileDialog.getSaveFileName(self, "Save Player Template", "players_template.xlsx", "Excel Files (*.xlsx)")
        if not path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Players"
            ws.append(["Name", "Position", "Base Price", "Category", "Notes", "Wishlist", "Max Bid"])
            ws.append(["Example Player", "Forward", 2000, "A", "", "Yes", 12000])
            wb.save(path)
            QMessageBox.information(self, "Template Created", "Excel template created successfully. This template is optional; mapped imports can use different headers.")
        except Exception as e:
            QMessageBox.critical(self, "Template Error", str(e))

    def export_excel(self):
        d = self.app.model
        default = f"{d.meta.get('name','auction').replace(' ','_')}_results.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "Export Auction Results", default, "Excel Files (*.xlsx)")
        if not path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Players"
            ws.append(["Player", "Position", "Category", "Base Price", "Status", "Team", "Sold Price", "Wishlist", "Max Bid"])
            for p in d.players:
                s = d.player_state(p["id"])
                ws.append([
                    p["name"],
                    p.get("position", ""),
                    p.get("category", ""),
                    p.get("base_price", 0),
                    s["status"],
                    s.get("team", ""),
                    s.get("price", 0) if s["status"] == "SOLD" else "",
                    "Yes" if p.get("wishlist") else "",
                    p.get("max_bid", 0),
                ])

            ts = wb.create_sheet("Teams")
            ts.append(["Team", "Players", "Spent", "Remaining", "Slots Left", "Safe Max Bid"])
            for t in d.teams:
                n = t["name"]
                s = d.team_summary(n)
                ts.append([n, s["players"], s["spent"], s["remaining"], s["slots"], d.safe_max_bid(n)])

            hs = wb.create_sheet("History")
            hs.append(["Time", "Player", "Status", "Team", "Price"])
            for e in d.events:
                p = d.get_player(e["player_id"])
                hs.append([
                    e.get("time", ""),
                    p["name"] if p else "Unknown",
                    "SOLD" if e["type"] == "sale" else "UNSOLD",
                    e.get("team", ""),
                    e.get("price", "") if e["type"] == "sale" else "",
                ])

            for sheet in wb.worksheets:
                sheet.freeze_panes = "A2"
                for cell in sheet[1]:
                    cell.font = cell.font.copy(bold=True)
                for cells in sheet.columns:
                    ml = max(len(str(c.value or "")) for c in cells)
                    sheet.column_dimensions[cells[0].column_letter].width = min(max(ml + 2, 12), 35)
            wb.save(path)
            QMessageBox.information(self, "Export Complete", "Auction results exported successfully.")
        except Exception as e:
            QMessageBox.critical(self, "Export Failed", str(e))


class SettingsPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 22, 24, 22)
        root.setSpacing(14)
        title = QLabel("Settings")
        title.setObjectName("SectionTitle")
        root.addWidget(title)

        box = QGroupBox("Auction Settings")
        form = QFormLayout(box)
        self.auction_name = QLineEdit()
        self.currency = QLineEdit()
        self.budget = QDoubleSpinBox(); self.budget.setRange(0, 1_000_000_000_000); self.budget.setDecimals(0)
        self.max_players = QSpinBox(); self.max_players.setRange(1, 100)
        self.my_team = QComboBox()
        self.theme = QComboBox(); self.theme.addItems(["Dark", "Light"])
        self.backup_keep = QSpinBox(); self.backup_keep.setRange(1, 100)
        self.common_enabled = QCheckBox("Use one base price for all players")
        self.common_price = QDoubleSpinBox(); self.common_price.setRange(0, 1_000_000_000_000); self.common_price.setDecimals(0); self.common_price.setSingleStep(500)
        self.common_enabled.toggled.connect(self.common_price.setEnabled)
        self.price_display = QComboBox(); self.price_display.addItem("Full numbers (100,000)", "full"); self.price_display.addItem("Compact K/M/B (100K, 1M)", "compact")
        self.quick_profile = QComboBox(); self.quick_profile.addItems(["Fine", "Standard", "Thousands", "Lakhs", "Millions", "Huge"])
        for label, w in [
            ("Auction name", self.auction_name), ("Currency", self.currency), ("Budget per team", self.budget),
            ("Max players per team", self.max_players), ("My Team", self.my_team), ("Theme", self.theme),
            ("Backups to keep", self.backup_keep), ("Price display", self.price_display), ("Quick price profile", self.quick_profile),
            ("Common base price", self.common_enabled), ("Common base amount", self.common_price)
        ]:
            form.addRow(label, w)
        root.addWidget(box)

        save = QPushButton("Apply Settings")
        save.setObjectName("Primary")
        save.clicked.connect(self.apply)
        root.addWidget(save, 0, Qt.AlignLeft)
        root.addStretch()

    def refresh(self):
        d = self.app.model
        self.auction_name.setText(d.meta.get("name", ""))
        self.currency.setText(d.meta.get("currency", "৳"))
        self.budget.setValue(float(d.meta.get("budget_per_team", 0)))
        self.max_players.setValue(int(d.meta.get("max_players_per_team", 12)))
        self.my_team.clear(); self.my_team.addItems([t["name"] for t in d.teams]); self.my_team.setCurrentText(d.meta.get("my_team", ""))
        self.theme.setCurrentText("Dark" if d.meta.get("theme", "dark") == "dark" else "Light")
        self.backup_keep.setValue(int(d.meta.get("backup_keep", 10)))
        mode = d.meta.get("price_display_mode", "full")
        idx = self.price_display.findData(mode); self.price_display.setCurrentIndex(idx if idx >= 0 else 0)
        profile = d.meta.get("quick_price_profile", "Standard")
        self.quick_profile.setCurrentText(profile if profile in LiveTrackerPage.QUICK_PRICE_PROFILES else "Standard")
        self.common_enabled.setChecked(bool(d.meta.get("use_common_base_price", False)))
        self.common_price.setValue(float(d.meta.get("common_base_price", 0) or 0))
        self.common_price.setEnabled(self.common_enabled.isChecked())

    def apply(self):
        d = self.app.model
        new_budget = self.budget.value()
        max_spent = max([d.team_summary(t["name"])["spent"] for t in d.teams], default=0)
        if new_budget < max_spent:
            QMessageBox.warning(self, "Invalid Budget", f"Budget cannot be lower than an already-spent total ({money(max_spent)}).")
            return
        new_max = self.max_players.value()
        max_count = max([d.team_summary(t["name"])["players"] for t in d.teams], default=0)
        if new_max < max_count:
            QMessageBox.warning(self, "Invalid Player Limit", f"Player limit cannot be lower than an existing squad size ({max_count}).")
            return
        d.meta.update({
            "name": self.auction_name.text().strip() or "Untitled Auction",
            "currency": self.currency.text().strip() or "৳",
            "budget_per_team": new_budget,
            "max_players_per_team": new_max,
            "my_team": self.my_team.currentText(),
            "theme": "dark" if self.theme.currentText() == "Dark" else "light",
            "backup_keep": self.backup_keep.value(),
            "price_display_mode": self.price_display.currentData() or "full",
            "quick_price_profile": self.quick_profile.currentText() or "Standard",
        })
        set_price_display_mode(d.meta.get("price_display_mode", "full"))
        d.set_common_base_price(self.common_enabled.isChecked(), self.common_price.value(), apply_existing=self.common_enabled.isChecked())
        self.app.apply_theme()
        self.app.after_change(auto_save=True)

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.model = AuctionData()
        set_price_display_mode(self.model.meta.get("price_display_mode", "full"))
        self.setWindowTitle(APP_NAME)
        self.setWindowIcon(QIcon(str(resource_path("assets", "app_icon.png"))))
        self.resize(1280, 800)
        self.setMinimumSize(1040, 680)

        self.build_menu()
        self.build_ui()
        QShortcut(QKeySequence("Ctrl+B"), self, activated=self.toggle_sidebar)
        self.apply_theme()
        for table in self.findChildren(QTableWidget):
            polish_table(table)

        # Start with a ready blank project, then ask user to create/open one
        self.new_auction(first_run=True)

    def icon(self, name):
        path = resource_path("assets", "icons", f"{name}.svg")
        return QIcon(str(path)) if path.exists() else QIcon()

    def build_menu(self):
        menu = self.menuBar()
        file_menu = menu.addMenu("File")

        new_action = QAction("New Auction", self)
        new_action.setShortcut("Ctrl+N")
        new_action.triggered.connect(self.new_auction)
        file_menu.addAction(new_action)

        open_action = QAction("Open Auction...", self)
        open_action.setShortcut("Ctrl+O")
        open_action.triggered.connect(self.open_auction)
        file_menu.addAction(open_action)

        save_action = QAction("Save", self)
        save_action.setShortcut("Ctrl+S")
        save_action.triggered.connect(self.save)
        file_menu.addAction(save_action)

        save_as_action = QAction("Save As...", self)
        save_as_action.setShortcut("Ctrl+Shift+S")
        save_as_action.triggered.connect(self.save_as)
        file_menu.addAction(save_as_action)

        undo_action = QAction("Undo Last Auction Entry", self)
        undo_action.setShortcut("Ctrl+Z")
        undo_action.triggered.connect(self.undo_last_global)
        file_menu.addAction(undo_action)

        file_menu.addSeparator()
        exit_action = QAction("Exit", self)
        exit_action.triggered.connect(self.close)
        file_menu.addAction(exit_action)

    def build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        outer = QHBoxLayout(central)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        self.sidebar = QFrame()
        self.sidebar.setObjectName("Sidebar")
        self.sidebar.setFixedWidth(220)
        sl = QVBoxLayout(self.sidebar)
        sl.setContentsMargins(10, 16, 10, 14)
        sl.setSpacing(8)

        brand_row = QHBoxLayout()
        brand_row.setContentsMargins(8, 2, 4, 4)
        logo = QLabel()
        logo.setPixmap(QIcon(str(resource_path("assets", "app_icon.png"))).pixmap(34, 34))
        logo.setFixedSize(36, 36)
        brand_row.addWidget(logo)
        brand_text = QVBoxLayout()
        brand_text.setSpacing(0)
        brand = QLabel("AUCTION TRACKER")
        brand.setObjectName("Brand")
        sub = QLabel("Simple. Fast. Clear.")
        sub.setObjectName("Muted")
        brand_text.addWidget(brand)
        brand_text.addWidget(sub)
        brand_row.addLayout(brand_text, 1)
        sl.addLayout(brand_row)

        self.project_label = QLabel("No auction")
        self.project_label.setObjectName("Muted")
        self.project_label.setWordWrap(True)
        self.project_label.setContentsMargins(10, 0, 8, 10)
        sl.addWidget(self.project_label)

        self.nav = QListWidget()
        nav_items = [
            ("Overview", "overview"),
            ("Live Tracker", "live"),
            ("Players", "players"),
            ("Teams", "teams"),
            ("My Team", "myteam"),
            ("Wishlist", "wishlist"),
            ("History", "history"),
            ("Import / Export", "import"),
            ("Settings", "settings"),
        ]
        self.nav.setIconSize(QSize(18, 18))
        for text, icon_name in nav_items:
            self.nav.addItem(QListWidgetItem(self.icon(icon_name), text))
        self.nav.currentRowChanged.connect(self.change_page)
        sl.addWidget(self.nav, 1)

        self.theme_button = QPushButton("Toggle Theme")
        self.theme_button.setIcon(self.icon("settings"))
        self.theme_button.clicked.connect(self.toggle_theme)
        sl.addWidget(self.theme_button)

        self.footer = QLabel()
        self.footer.setObjectName("Muted")
        self.footer.setWordWrap(True)
        self.footer.setOpenExternalLinks(True)
        self.footer.setTextInteractionFlags(Qt.TextBrowserInteraction)
        self.footer.setContentsMargins(8, 8, 8, 2)
        sl.addWidget(self.footer)

        outer.addWidget(self.sidebar)

        content = QVBoxLayout()
        content.setContentsMargins(0, 0, 0, 0)
        content.setSpacing(0)

        top = QFrame()
        top.setObjectName("TopBar")
        top.setFixedHeight(58)
        tl = QHBoxLayout(top)
        tl.setContentsMargins(12, 10, 18, 10)
        self.sidebar_toggle = QPushButton("☰")
        self.sidebar_toggle.setToolTip("Hide / show sidebar (Ctrl+B)")
        self.sidebar_toggle.setFixedSize(40, 36)
        self.sidebar_toggle.clicked.connect(self.toggle_sidebar)
        tl.addWidget(self.sidebar_toggle)
        self.page_title = QLabel("Overview")
        self.page_title.setObjectName("Brand")
        tl.addWidget(self.page_title)
        tl.addStretch()
        self.save_status = QLabel("Unsaved")
        self.save_status.setObjectName("Muted")
        tl.addWidget(self.save_status)
        new_btn = QPushButton("New")
        new_btn.setIcon(self.icon("plus"))
        new_btn.clicked.connect(self.new_auction)
        open_btn = QPushButton("Open")
        open_btn.setIcon(self.icon("folder"))
        open_btn.clicked.connect(self.open_auction)
        save_btn = QPushButton("Save")
        save_btn.setIcon(self.icon("save"))
        save_btn.setObjectName("Primary")
        save_btn.clicked.connect(self.save)
        tl.addWidget(new_btn)
        tl.addWidget(open_btn)
        tl.addWidget(save_btn)
        content.addWidget(top)

        self.stack = QStackedWidget()
        self.pages = [
            OverviewPage(self),
            LiveTrackerPage(self),
            PlayersPage(self),
            TeamsPage(self),
            MyTeamPage(self),
            WishlistPage(self),
            HistoryPage(self),
            ImportExportPage(self),
            SettingsPage(self),
        ]
        for p in self.pages:
            self.stack.addWidget(p)
        content.addWidget(self.stack, 1)

        wrapper = QWidget()
        wrapper.setLayout(content)
        outer.addWidget(wrapper, 1)

        self.nav.setCurrentRow(0)

    def toggle_sidebar(self):
        currently_visible = self.sidebar.isVisible()
        self.sidebar.setVisible(not currently_visible)
        if currently_visible:
            self.statusBar().showMessage("Sidebar hidden. Press Ctrl+B or click ☰ to show it again.", 2200)
        else:
            self.statusBar().showMessage("Sidebar shown.", 1500)

    def change_page(self, idx):
        if idx < 0:
            return
        self.stack.setCurrentIndex(idx)
        item = self.nav.item(idx)
        self.page_title.setText(item.text() if item else "")
        self.refresh_page(idx)

    def refresh_page(self, idx):
        page = self.pages[idx]
        if hasattr(page, "refresh"):
            page.refresh()

    def refresh_all(self):
        set_price_display_mode(self.model.meta.get("price_display_mode", "full"))
        self.project_label.setText(
            f"{self.model.meta.get('name','')}\n"
            f"My Team: {self.model.meta.get('my_team','—')}"
        )
        self.update_footer()
        self.setWindowTitle(
            f"{self.model.meta.get('name','Untitled')} — {APP_NAME}"
            + (f" — {self.model.file_path.name}" if self.model.file_path else "")
        )
        if hasattr(self, "save_status"):
            self.save_status.setText("Saved" if self.model.file_path else "Unsaved")
        for p in self.pages:
            if hasattr(p, "refresh"):
                p.refresh()

    def update_footer(self):
        year = datetime.now().year
        links = (
            f'<a href="{FACEBOOK_URL}" style="text-decoration:none;">Facebook</a>'
            f' &nbsp;•&nbsp; <a href="{INSTAGRAM_URL}" style="text-decoration:none;">Instagram</a>'
            f' &nbsp;•&nbsp; <a href="{GITHUB_URL}" style="text-decoration:none;">GitHub</a>'
        )
        self.footer.setText(
            f"<div style='font-size:11px; line-height:1.5;'>© {year} Auction Tracker Pro<br>Made by {DEVELOPER_NAME}<br>{links}</div>"
        )

    def apply_theme(self):
        theme = self.model.meta.get("theme", "dark")
        QApplication.instance().setStyleSheet((DARK_QSS if theme == "dark" else LIGHT_QSS) + COMMON_QSS)

    def toggle_theme(self):
        self.model.meta["theme"] = "light" if self.model.meta.get("theme","dark") == "dark" else "dark"
        self.apply_theme()
        self.after_change(auto_save=True)

    def new_auction(self, checked=False, first_run=False):
        dlg = NewAuctionDialog(self)
        if dlg.exec():
            v = dlg.values()
            current_theme = self.model.meta.get("theme", "dark")
            self.model.new(v["name"], v["teams"], v["budget"], v["max_players"], v["my_team"], v["currency"], v.get("use_common_base_price", False), v.get("common_base_price", 0))
            self.model.meta["theme"] = current_theme
            self.apply_theme()
            self.refresh_all()
            self.nav.setCurrentRow(2)
        elif first_run:
            self.refresh_all()

    def open_auction(self):
        path, _ = QFileDialog.getOpenFileName(self, "Open Auction", "", "Auction Files (*.json)")
        if not path:
            return
        try:
            self.model.load(path)
            self.apply_theme()
            self.refresh_all()
            self.nav.setCurrentRow(0)
        except Exception as e:
            QMessageBox.critical(self, "Open Failed", str(e))

    def save(self):
        if not self.model.file_path:
            return self.save_as()
        try:
            self.model.save()
            self.backup_project()
            self.refresh_all()
            self.save_status.setText("Saved")
            self.statusBar().showMessage("Auction saved and backed up.", 2500)
            return True
        except Exception as e:
            QMessageBox.critical(self, "Save Failed", str(e))
            return False

    def save_as(self):
        default = f"{self.model.meta.get('name','auction').replace(' ','_')}.json"
        path, _ = QFileDialog.getSaveFileName(self, "Save Auction", default, "Auction Files (*.json)")
        if not path:
            return False
        if not path.lower().endswith(".json"):
            path += ".json"
        try:
            self.model.save(path)
            self.backup_project()
            self.refresh_all()
            self.save_status.setText("Saved")
            self.statusBar().showMessage("Auction saved and backup folder created.", 2500)
            return True
        except Exception as e:
            QMessageBox.critical(self, "Save Failed", str(e))
            return False

    def backup_project(self):
        if not self.model.file_path:
            return
        try:
            keep = max(1, int(self.model.meta.get("backup_keep", 10)))
            backup_dir = self.model.file_path.parent / "AuctionTrackerBackups" / self.model.file_path.stem
            backup_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            target = backup_dir / f"{self.model.file_path.stem}_{stamp}.json"
            shutil.copy2(self.model.file_path, target)
            files = sorted(backup_dir.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True)
            for old in files[keep:]:
                old.unlink(missing_ok=True)
        except Exception as e:
            self.statusBar().showMessage(f"Backup warning: {e}", 3500)

    def undo_last_global(self):
        e = self.model.undo_last()
        if not e:
            self.statusBar().showMessage("Nothing to undo.", 1800)
            return
        self.after_change(auto_save=True)

    def after_change(self, auto_save=False):
        self.refresh_all()
        if hasattr(self, "save_status"):
            self.save_status.setText("Unsaved")
        if auto_save and self.model.file_path:
            try:
                self.model.save()
                self.backup_project()
                if hasattr(self, "save_status"):
                    self.save_status.setText("Saved")
                self.statusBar().showMessage("Saved automatically. Backup updated.", 1800)
            except Exception as e:
                if hasattr(self, "save_status"):
                    self.save_status.setText("Save failed")
                self.statusBar().showMessage(f"Auto-save failed: {e}", 3500)

    def closeEvent(self, event):
        if self.model.file_path:
            try:
                self.model.save()
            except Exception:
                pass
        event.accept()


def main():
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setApplicationVersion(APP_VERSION)
    app.setWindowIcon(QIcon(str(resource_path("assets", "app_icon.png"))))
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
